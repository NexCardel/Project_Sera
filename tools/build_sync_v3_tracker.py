"""
tools/build_sync_v3_tracker.py
------------------------------
Builds the Sera Sync v3 tracker from docs/sera-sync-v3-blueprint.md:

  docs/sera-sync-v3-plan.json    fixed data: WPs (tier, dependencies, review), models, phase checks
  docs/sera-sync-v3-status.csv   live WP status   (written only by tools/sync_v3_tracker.py)
  docs/sera-sync-v3-checks.csv   phase-check results (same)
  docs/sera-sync-v3-agents.xlsx  read-only viewer; pulls the two CSVs in with Power Query

The CSVs are never overwritten if they already exist, so rebuilding keeps all progress.
Run it again whenever the blueprint's §3 table changes:

  venv\\Scripts\\python tools\\build_sync_v3_tracker.py
  venv\\Scripts\\python tools\\build_sync_v3_tracker.py --no-excel     (skip the Power Query step; tests use this)
  venv\\Scripts\\python tools\\build_sync_v3_tracker.py --migrate-from old.xlsx   (one-off: copy statuses out of the old tracker)

Close the viewer in Excel before rebuilding it.
"""

import argparse
import csv
import json
import os
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
DOCS = REPO / "docs"

STATUS_FIELDS = ["WP", "Status", "Model used", "Reviewed by", "Commit", "Notes", "Updated"]
CHECK_FIELDS = ["Check", "Result", "Date", "Notes"]

MODELS = [
    # name, provider, id, status, T1, T2, T3, can review T3 work, notes
    ("Claude Fable 5.1", "Anthropic", "claude-fable-5-1", "GA (Sep 2026)", "Yes", "Yes", "Yes", "Yes",
     "Most capable Claude model. First choice for T3 work and all reviews."),
    ("Claude Opus 5.5", "Anthropic", "claude-opus-5-5", "GA", "Yes", "Yes", "Yes", "Yes",
     "T3-capable. Alternative to Fable for T3 work and reviews."),
    ("Claude Sonnet 5", "Anthropic", "claude-sonnet-5", "GA", "Yes", "Yes", "No", "No", "Default for T2."),
    ("Claude Haiku 4.5", "Anthropic", "claude-haiku-4-5-20251001", "GA", "Yes", "No", "No", "No", "Default for T1."),
    ("Gemini 3.8 Flash", "Google", "gemini-3.8-flash", "Stable (GA 2026-09-02)", "Yes", "Yes", "No", "No",
     "Best Gemini choice for T2 (Google: long-horizon software engineering, agents)."),
    ("Gemini 3.1 Pro", "Google", "gemini-3.1-pro-preview", "Preview", "Yes", "Yes", "No", "No",
     "Strong reasoning; preview model, so behaviour can change between sessions."),
    ("Gemini 3.7 Flash", "Google", "gemini-3.7-flash", "Stable", "Yes", "Yes", "No", "No",
     "Previous-generation Flash for complex coding; fine for T2."),
    ("Gemini 3.6 Flash", "Google", "gemini-3.6-flash", "Stable", "Yes", "Caution", "No", "No",
     "T2 only with a Fable/Opus review of the diff."),
    ("Gemini 3 Flash", "Google", "gemini-3-flash-preview", "Preview", "Yes", "Caution", "No", "No",
     "Older preview; T2 only with a Fable/Opus review of the diff."),
    ("Gemini 2.5 Pro", "Google", "gemini-2.5-pro", "Legacy (access-limited)", "Yes", "Caution", "No", "No",
     "Only if your account still has access; T2 only with review."),
    ("Gemini 3.5 Flash", "Google", "gemini-3.5-flash", "Stable (legacy)", "Yes", "No", "No", "No", "Baseline model; T1 only."),
    ("Gemini 3.5 Flash-Lite", "Google", "gemini-3.5-flash-lite", "Stable", "Yes", "No", "No", "No", "Fast/cheap; T1 only."),
    ("Gemini 3.1 Flash-Lite", "Google", "gemini-3.1-flash-lite", "Stable", "Yes", "No", "No", "No", "Fast/cheap; T1 only."),
    ("Gemini 2.5 Flash", "Google", "gemini-2.5-flash", "Legacy (access-limited)", "Yes", "No", "No", "No", "T1 only."),
    ("Gemini 2.5 Flash-Lite", "Google", "gemini-2.5-flash-lite", "Legacy (access-limited)", "Yes", "No", "No", "No", "T1 only."),
]

TIERS = [
    ("T1", "Mechanical", "Claude Haiku 4.5", "Any Gemini model on the Models sheet"),
    ("T2", "Standard", "Claude Sonnet 5",
     "Gemini 3.8 Flash, 3.1 Pro, 3.7 Flash (3.6 Flash / 3 Flash / 2.5 Pro only with review)"),
    ("T3", "Expert", "Claude Fable 5.1 or Opus 5.5", "Not allowed (owner decision)"),
]

PHASES = [(0, "Hotfix (2.x)"), (1, "Office key (3.0)"), (2, "Devices & pairing (3.0)"),
          (3, "Change replication (3.1 shadow / 3.2 live)"), (4, "Cleanup (3.3)")]

PHASE_CHECKS = [
    (0, "Installer adds the firewall rule: netsh advfirewall firewall show rule name=\"Amas Sera Sync\"", "P0-9a"),
    (0, "Sera Sync panel warns when the PC's network is set to Public", "P0-9b"),
    (0, "Spare PC: first run → Join office → approve on the other PC → enter office password → app opens", "P0-6"),
    (0, "Manual push from the dialog lands only after the receiving PC restarts, and the DB opens", "P0-4"),
    (0, "A Wi-Fi PC and a wired PC see each other after Add PC by IP", "P0-10"),
    (0, "Editing a client no longer triggers 'Live update synced' toasts on other PCs", "P0-1"),
    (1, "Admin PC: Convert to office key succeeds; app opens without a password prompt", "P1-4"),
    (1, "Recovery kit exported to USB and stored safely", "P1-6"),
    (1, "Log into a second Windows account on the admin PC → recovery dialog → master password unlocks", "P1-6"),
    (1, "DOM Parser and SDC Parser still run on the converted PC", "P1-3"),
    (2, "Add workstation → code → Join office on another PC works end to end", "P2-4, P2-6"),
    (2, "Wrong code 3 times closes the pairing window", "P2-4"),
    (2, "Removed PC can no longer sync (panel shows it as removed)", "P2-7"),
    (2, "Become admin with the master password on a second PC; the old admin PC shows it is no longer admin", "P2-2"),
    (2, "Every legacy PC: Rejoin → salvage dry run reviewed → import; salvage report saved", "P2-8"),
    (3, "Shadow mode on all PCs for 7 days; zero capture mismatches; replica digests converge", "P3-7"),
    (3, "Two people edit different fields of the same client on two PCs → both edits appear everywhere", "P3-4"),
    (3, "Set one PC's clock 2 hours ahead → panel shows the clock warning; set it back", "P3-3"),
    (3, "On a non-admin PC, enter admin mode and change a setting → it appears on every PC", "P3-4"),
    (3, "On a non-admin PC, try to edit the staff list → told to use the admin PC or Become admin", "P3-1"),
    (3, "Go-live restart on every PC; clients count matches on all PCs", "P3-9"),
    (4, "An old 2.x install can no longer push to the office", "P4-1"),
    (4, "A daily backup folder appears in backups/", "P4-3a"),
    (4, "Restore dry run shows sensible counts (cancel at the RESTORE prompt unless you mean it)", "P4-3b"),
]


# ------------------------------------------------------------------ plan (from the blueprint)

def parse_blueprint(doc_path: Path) -> list[dict]:
    text = doc_path.read_text(encoding="utf-8")
    sec3 = text.split("## 3. Phase plan at a glance", 1)[1].split("**Release grouping**", 1)[0]
    rows = []
    for line in sec3.splitlines():
        if not line.startswith("| P"):
            continue
        wp, phase, what, tier, deps, review = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append({"wp": wp, "phase": int(re.match(r"\d", phase).group(0)), "what": what.replace("`", ""),
                     "tier": tier.replace("*", ""), "deps_raw": deps.replace("*", ""),
                     "review_raw": review.replace("*", "")})
    if not rows:
        raise SystemExit("No WP rows found in the blueprint's §3 table.")
    ids = [r["wp"] for r in rows]
    by_phase = {p: [r["wp"] for r in rows if r["phase"] == p] for p in range(5)}
    wps = []
    for r in rows:
        raw = r["deps_raw"]
        if raw in ("–", "-", ""):
            deps = []
        elif raw == "P0 released":
            deps = by_phase[0]
        elif raw == "all":
            deps = [w for w in by_phase[r["phase"]] if w != r["wp"]]
        else:
            deps = re.findall(r"P\d-\d+[ab]?", raw)
        unknown = [d for d in deps if d not in ids]
        if unknown:
            raise SystemExit(f"{r['wp']} depends on unknown WP(s): {unknown}")
        note = ""
        if raw == "P0 released":
            note = "Phase 0 released to all PCs"
        elif "criteria" in raw:
            note = "Plus: P3-7 go-live criteria met (owner decides)"
        elif "all PCs" in raw:
            note = "Plus: go-live done on every PC"
        elif "order matters" in raw:
            note = "Must merge P0-1 first (see blueprint P0-2)"
        rv = r["review_raw"]
        review_text = "No" if rv in ("–", "-") else ("Yes – Fable 5.1 / Opus 5.5" if rv == "T3" else rv)
        wps.append({"wp": r["wp"], "phase": r["phase"], "what": r["what"], "tier": r["tier"], "deps": deps,
                    "dep_note": note, "review": review_text, "review_required": review_text.startswith("Yes")})
    return wps


def build_plan(doc_path: Path) -> dict:
    return {
        "source": "docs/sera-sync-v3-blueprint.md §3 (generated by tools/build_sync_v3_tracker.py; don't edit by hand)",
        "wps": parse_blueprint(doc_path),
        "models": [{"name": m[0], "provider": m[1], "id": m[2], "status": m[3], "T1": m[4], "T2": m[5],
                    "T3": m[6], "review": m[7], "notes": m[8]} for m in MODELS],
        "tiers": [{"tier": t[0], "meaning": t[1], "claude": t[2], "gemini": t[3]} for t in TIERS],
        "phases": [{"phase": p, "name": n} for p, n in PHASES],
        "phase_checks": [{"n": i + 1, "phase": p, "check": c, "wp": w} for i, (p, c, w) in enumerate(PHASE_CHECKS)],
    }


def write_text_atomic(path: Path, text: str):
    tmp = path.with_name(path.name + ".tmp")
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        f.write(text)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)


def write_csv_atomic(path: Path, fields: list, rows: list[dict]):
    import io
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields, lineterminator="\r\n")
    w.writeheader()
    for r in rows:
        w.writerow({k: r.get(k, "") or "" for k in fields})
    write_text_atomic(path, buf.getvalue())


def migrate_from_xlsx(old: Path) -> tuple[list[dict], list[dict]]:
    """Copies the input columns out of the pre-CSV tracker (status was typed into the xlsx)."""
    from openpyxl import load_workbook
    wb = load_workbook(old, data_only=True)
    ws = wb["Work Packages"]
    head = {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1)}
    statuses = []
    for r in range(5, ws.max_row + 1):
        wp = ws.cell(row=r, column=head["WP"]).value
        if not wp:
            continue
        get = lambda name: ws.cell(row=r, column=head[name]).value if name in head else None
        statuses.append({"WP": wp, "Status": get("Status") or "Not started", "Model used": get("Model used") or "",
                         "Reviewed by": get("Reviewed by") or "", "Commit": get("Commit") or "",
                         "Notes": get("Notes") or "", "Updated": ""})
    checks = []
    pc = wb["Phase Checks"]
    phead = {pc.cell(row=3, column=c).value: c for c in range(1, pc.max_column + 1)}
    for r in range(4, pc.max_row + 1):
        text = pc.cell(row=r, column=phead["Check"]).value
        if not text:
            continue
        n = next((c["n"] for c in build_plan_checks() if c["check"] == text), None)
        if n is None:
            continue
        date = pc.cell(row=r, column=phead["Date"]).value
        checks.append({"Check": n, "Result": pc.cell(row=r, column=phead["Result"]).value or "Not run",
                       "Date": str(date) if date else "", "Notes": pc.cell(row=r, column=phead["Notes"]).value or ""})
    return statuses, checks


def build_plan_checks():
    return [{"n": i + 1, "check": c} for i, (_p, c, _w) in enumerate(PHASE_CHECKS)]


# ------------------------------------------------------------------ viewer workbook

def build_viewer(plan: dict, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.comments import Comment

    FONT = "Arial"
    F_BODY, F_BOLD = Font(name=FONT, size=10), Font(name=FONT, size=10, bold=True)
    F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    F_TITLE = Font(name=FONT, size=14, bold=True)
    F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
    FILL_HEAD = PatternFill("solid", fgColor="2F4F6F")
    FILL_LIVE = PatternFill("solid", fgColor="EDEDED")      # comes from the status CSV
    FILL_T = {"T1": PatternFill("solid", fgColor="E2EFDA"), "T2": PatternFill("solid", fgColor="DDEBF7"),
              "T3": PatternFill("solid", fgColor="FCE4D6")}
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    WRAP = Alignment(wrap_text=True, vertical="top")
    green, grey = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="D9D9D9")
    red, amber = PatternFill("solid", fgColor="FFC7CE"), PatternFill("solid", fgColor="FFEB9C")

    def header(ws, row, labels, widths):
        for i, lab in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=lab)
            c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[c.column_letter].width = widths[i - 1]

    def body(c, bold=False):
        c.font, c.alignment, c.border = (F_BOLD if bold else F_BODY), WRAP, BORDER

    wb = Workbook()

    # Models
    ms = wb.active
    ms.title = "Models"
    ms["A1"], ms["A1"].font = "Model × tier matrix", F_TITLE
    ms["A2"] = ("Yes = allowed. Caution = allowed only if a Fable 5.1 / Opus 5.5 session reviews the diff before merge. "
                "No = don't use for that tier. Gemini is never used for T3 (owner decision 2026-09-23).")
    ms["A2"].font = F_NOTE
    header(ms, 4, ["Model", "Provider", "Model ID", "Status", "T1", "T2", "T3", "Can review T3 work", "Notes"],
           [24, 11, 28, 22, 8, 10, 8, 12, 60])
    for i, m in enumerate(plan["models"]):
        for col, key in enumerate(["name", "provider", "id", "status", "T1", "T2", "T3", "review", "notes"], start=1):
            body(ms.cell(row=5 + i, column=col, value=m[key]), bold=(col == 1))
    m_first, m_last = 5, 4 + len(plan["models"])
    r = m_last + 2
    ms.cell(row=r, column=1, value="Sources").font = F_BOLD
    ms.cell(row=r + 1, column=1, value=(
        "Gemini list: Google AI for Developers, 'Models | Gemini API', https://ai.google.dev/gemini-api/docs/models "
        "(fetched 2026-09-23). Claude list: Claude Code model list + Anthropic's Fable 5.1 announcement. "
        "Tier permissions are Claude's recommendation for this project, not vendor ratings.")).font = F_NOTE
    tt = r + 4
    ms.cell(row=tt - 1, column=1, value="Tier defaults").font = F_BOLD
    for i, lab in enumerate(["Tier", "Meaning", "Recommended Claude", "Gemini allowed"], start=1):
        c = ms.cell(row=tt, column=i, value=lab)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
    for i, t in enumerate(plan["tiers"]):
        for col, key in enumerate(["tier", "meaning", "claude", "gemini"], start=1):
            body(ms.cell(row=tt + 1 + i, column=col, value=t[key]), bold=(col == 1))
    t_first, t_last = tt + 1, tt + len(plan["tiers"])

    # Work Packages
    ws = wb.create_sheet("Work Packages", 0)
    ws["A1"], ws["A1"].font = "Sera Sync v3 — work packages", F_TITLE
    ws["A2"] = ("Grey columns come from docs/sera-sync-v3-status.csv, which agents update with tools/sync_v3_tracker.py. "
                "Data → Refresh All (Ctrl+Alt+F5) to see the latest; it also refreshes on open and every minute. "
                "Don't type into this sheet: a refresh overwrites it.")
    ws["A2"].font = F_NOTE
    H, first = 4, 5
    last = first + len(plan["wps"]) - 1
    header(ws, H, ["WP", "Phase", "What", "Tier", "Depends on", "Dependency note", "Review before merge",
                   "Recommended Claude", "Gemini allowed", "Status", "Ready?", "Unfinished deps", "Model used",
                   "Model allowed?", "Reviewed by", "Commit", "Notes", "helper"],
           [8, 7, 46, 6, 22, 26, 22, 22, 30, 13, 14, 10, 22, 16, 18, 12, 40, 6])
    rng = lambda col: f"${col}${first}:${col}${last}"

    def live(n, idx, default=""):
        look = f"INDEX(StatusData!${idx}:${idx},MATCH($A{n},StatusData!$A:$A,0))"
        if default:
            return f'=IFERROR(IF({look}="","{default}",""&{look}),"{default}")'
        return f'=IFERROR(""&{look},"")'

    for i, wp in enumerate(plan["wps"]):
        n = first + i
        vals = {
            "A": wp["wp"], "B": wp["phase"], "C": wp["what"], "D": wp["tier"], "E": ",".join(wp["deps"]),
            "F": wp["dep_note"], "G": wp["review"],
            "H": f"=INDEX(Models!$C${t_first}:$C${t_last},MATCH(D{n},Models!$A${t_first}:$A${t_last},0))",
            "I": f"=INDEX(Models!$D${t_first}:$D${t_last},MATCH(D{n},Models!$A${t_first}:$A${t_last},0))",
            "J": live(n, "B", "Not started"),
            "K": f'=IF(J{n}="Done","Done",IF(L{n}>0,"Waiting on "&L{n},IF(J{n}="Not started","Ready",J{n})))',
            "L": (f'=IF(E{n}="",0,SUMPRODUCT(ISNUMBER(SEARCH(","&{rng("A")}&",",'
                  f'","&SUBSTITUTE(E{n}," ","")&","))*({rng("J")}<>"Done")))'),
            "M": live(n, "C"),
            "N": (f'=IF(M{n}="","",IFERROR(IF(INDEX(Models!$E${m_first}:$G${m_last},'
                  f'MATCH(M{n},Models!$A${m_first}:$A${m_last},0),VALUE(RIGHT(D{n},1)))="Yes","OK",'
                  f'IF(INDEX(Models!$E${m_first}:$G${m_last},MATCH(M{n},Models!$A${m_first}:$A${m_last},0),'
                  f'VALUE(RIGHT(D{n},1)))="Caution","Needs T3 review","NOT ALLOWED")),"Unknown model"))'),
            "O": live(n, "D"), "P": live(n, "E"), "Q": live(n, "F"),
            "R": f'=IF(K{n}="Ready",B{n},"")',
        }
        for col, v in vals.items():
            c = ws[f"{col}{n}"]
            c.value = v
            body(c, bold=(col == "A"))
            if col in "JMOPQ":
                c.fill = FILL_LIVE
        ws[f"D{n}"].fill = FILL_T[wp["tier"]]
    ws.conditional_formatting.add(f"K{first}:K{last}", FormulaRule(formula=[f'K{first}="Ready"'], fill=green))
    ws.conditional_formatting.add(f"K{first}:K{last}", FormulaRule(formula=[f'K{first}="Done"'], fill=grey))
    ws.conditional_formatting.add(f"N{first}:N{last}", FormulaRule(formula=[f'N{first}="NOT ALLOWED"'], fill=red))
    ws.conditional_formatting.add(f"N{first}:N{last}", FormulaRule(formula=[f'N{first}="Needs T3 review"'], fill=amber))
    ws.conditional_formatting.add(f"N{first}:N{last}", FormulaRule(formula=[f'N{first}="OK"'], fill=green))
    ws.conditional_formatting.add(f"A{first}:Q{last}", FormulaRule(formula=[f'$J{first}="Blocked"'], fill=red))
    ws.freeze_panes = ws.cell(row=first, column=4)
    ws.auto_filter.ref = f"A{H}:Q{last}"
    ws.column_dimensions["R"].hidden = True
    ws["E4"].comment = Comment("WP ids separated by commas. 'Unfinished deps' counts those not yet Done.", "Sera")

    # Summary
    ss = wb.create_sheet("Summary", 1)
    ss["A1"], ss["A1"].font = "Progress summary", F_TITLE
    header(ss, 3, ["Phase", "Name", "WPs", "Done", "In progress / review", "Blocked", "% done", "Next ready WP"],
           [7, 40, 7, 7, 12, 9, 9, 14])
    W = "'Work Packages'"
    for i, ph in enumerate(plan["phases"]):
        n = 4 + i
        vals = [ph["phase"], ph["name"], f"=COUNTIFS({W}!{rng('B')},A{n})",
                f'=COUNTIFS({W}!{rng("B")},A{n},{W}!{rng("J")},"Done")',
                f'=COUNTIFS({W}!{rng("B")},A{n},{W}!{rng("J")},"In progress")'
                f'+COUNTIFS({W}!{rng("B")},A{n},{W}!{rng("J")},"In review")',
                f'=COUNTIFS({W}!{rng("B")},A{n},{W}!{rng("J")},"Blocked")',
                f"=IF(C{n}=0,0,D{n}/C{n})",
                f'=IFERROR(INDEX({W}!{rng("A")},MATCH(A{n},{W}!{rng("R")},0)),"none")']
        for col, v in enumerate(vals, start=1):
            body(ss.cell(row=n, column=col, value=v), bold=(col == 1))
        ss.cell(row=n, column=7).number_format = "0%"
    tot = 4 + len(plan["phases"])
    ss.cell(row=tot, column=2, value="Total").font = F_BOLD
    for col in "CDEF":
        ss[f"{col}{tot}"] = f"=SUM({col}4:{col}{tot - 1})"
        body(ss[f"{col}{tot}"], bold=True)
    ss[f"G{tot}"] = f"=IF(C{tot}=0,0,D{tot}/C{tot})"
    body(ss[f"G{tot}"], bold=True)
    ss[f"G{tot}"].number_format = "0%"
    t0 = tot + 3
    ss.cell(row=t0 - 1, column=1, value="By tier").font = F_BOLD
    for i, lab in enumerate(["Tier", "Meaning", "WPs", "Done", "Need review"], start=1):
        c = ss.cell(row=t0, column=i, value=lab)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
    for i, t in enumerate(plan["tiers"]):
        n = t0 + 1 + i
        for col, v in enumerate([t["tier"], t["meaning"], f"=COUNTIFS({W}!{rng('D')},A{n})",
                                 f'=COUNTIFS({W}!{rng("D")},A{n},{W}!{rng("J")},"Done")',
                                 f'=COUNTIFS({W}!{rng("D")},A{n},{W}!{rng("G")},"Yes*")'], start=1):
            body(ss.cell(row=n, column=col, value=v), bold=(col == 1))
    ss.cell(row=t0 + 5, column=1, value="Status last refreshed from the CSV:").font = F_NOTE
    ss.cell(row=t0 + 5, column=4, value='=IFERROR(MAX(StatusData!G:G)&"","")').font = F_NOTE

    # Phase Checks
    pcs = wb.create_sheet("Phase Checks", 2)
    pcs["A1"], pcs["A1"].font = "Hands-on checks before releasing each phase (blueprint §8.4)", F_TITLE
    pcs["A2"] = ("Record results with: venv\\Scripts\\python tools\\sync_v3_tracker.py check <#> --result Pass "
                 "(or ask an agent). Grey columns come from docs/sera-sync-v3-checks.csv.")
    pcs["A2"].font = F_NOTE
    header(pcs, 3, ["#", "Phase", "Check", "Related WP", "Result", "Date", "Notes"], [5, 7, 80, 12, 11, 12, 40])
    for i, chk in enumerate(plan["phase_checks"]):
        n = 4 + i
        # ""& because column A holds the check number as a number, but Power Query loads every
        # CSV column as text (AsText below) -- MATCH never matches 1 to "1".
        look = lambda idx: f"INDEX(ChecksData!${idx}:${idx},MATCH(\"\"&$A{n},ChecksData!$A:$A,0))"
        vals = [chk["n"], chk["phase"], chk["check"], chk["wp"],
                f'=IFERROR(IF({look("B")}="","Not run",""&{look("B")}),"Not run")',
                f'=IFERROR(""&{look("C")},"")', f'=IFERROR(""&{look("D")},"")']
        for col, v in enumerate(vals, start=1):
            c = pcs.cell(row=n, column=col, value=v)
            body(c)
            if col >= 5:
                c.fill = FILL_LIVE
    pl = 3 + len(plan["phase_checks"])
    pcs.conditional_formatting.add(f"E4:E{pl}", FormulaRule(formula=['E4="Pass"'], fill=green))
    pcs.conditional_formatting.add(f"E4:E{pl}", FormulaRule(formula=['E4="Fail"'], fill=red))
    pcs.freeze_panes = "C4"

    # How to use
    hs = wb.create_sheet("How to use", 0)
    hs.column_dimensions["A"].width = 110
    lines = [
        ("How to use this workbook", F_TITLE), ("", F_BODY),
        ("This workbook only displays status. Agents record progress in two small text files next to it", F_BOLD),
        ("(docs/sera-sync-v3-status.csv and docs/sera-sync-v3-checks.csv) using tools/sync_v3_tracker.py.", F_BOLD),
        ("You can keep this file open while agents work.", F_BOLD), ("", F_BODY),
        ("To see the latest status: Data → Refresh All (Ctrl+Alt+F5). It also refreshes when opened and every minute.", F_BODY),
        ("Don't type into the sheets: the next refresh overwrites it. To change a status, ask an agent or run the tool.", F_BODY),
        ("When Excel asks to save on close, either answer is fine — the status lives in the CSV files, not here.", F_BODY),
        ("", F_BODY),
        ("What the agents record:", F_BOLD),
        ("• Implementing agent: Status = In progress + Model used when it starts; Status = In review + Notes when it finishes.", F_BODY),
        ("• Reviewing agent (Fable 5.1 / Opus 5.5): Reviewed by.", F_BODY),
        ("• Committing agent: Status = Done + Commit. The tool refuses Done without a commit, or without a required review.", F_BODY),
        ("• The tool also refuses a model not allowed for the WP's tier, and starting a WP whose dependencies aren't Done.", F_BODY),
        ("", F_BODY),
        ("What you do:", F_BOLD),
        ("1. Pick a WP whose 'Ready?' is green (or Summary → Next ready WP).", F_BODY),
        ("2. Start a fresh chat with a model allowed for its tier and paste the prompt from blueprint §8.1.", F_BODY),
        ("3. After a phase: do the Phase Checks and record the results (tool command on that sheet, or ask an agent).", F_BODY),
        ("", F_BODY),
        ("Rebuild (only when the blueprint's §3 table changes; close this file first): "
         "venv\\Scripts\\python tools\\build_sync_v3_tracker.py — progress in the CSVs is kept.", F_NOTE),
        ("If refresh ever shows an error about a missing file, the project folder moved: run the rebuild command.", F_NOTE),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = hs.cell(row=i, column=1, value=t)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")

    # Placeholders the Power Query step fills (formulas already point at them)
    for name, fields in (("StatusData", STATUS_FIELDS), ("ChecksData", CHECK_FIELDS)):
        s = wb.create_sheet(name)
        for i, f in enumerate(fields, start=1):
            s.cell(row=1, column=i, value=f)
        s.sheet_state = "hidden"

    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(out)


def attach_power_query(xlsx: Path, status_csv: Path, checks_csv: Path):
    """Uses Excel itself (COM) to add the two CSV queries, refresh them and save."""
    import comtypes.client

    def m_formula(path: Path) -> str:
        p = str(path.resolve()).replace('"', '""')
        return ('let Source = Csv.Document(File.Contents("' + p + '"),[Delimiter=",", Encoding=65001, '
                'QuoteStyle=QuoteStyle.Csv]), Promoted = Table.PromoteHeaders(Source, [PromoteAllScalars=true]), '
                'AsText = Table.TransformColumnTypes(Promoted, List.Transform(Table.ColumnNames(Promoted), '
                'each {_, type text})) in AsText')

    xl = comtypes.client.CreateObject("Excel.Application")
    try:
        xl.Visible = False
        xl.DisplayAlerts = False
        wb = xl.Workbooks.Open(str(xlsx.resolve()))
        for sheet_name, query, path in (("StatusData", "SeraStatusCsv", status_csv),
                                        ("ChecksData", "SeraChecksCsv", checks_csv)):
            wb.Queries.Add(query, m_formula(path))
            ws = wb.Worksheets(sheet_name)
            ws.Cells.Clear()
            lo = ws.ListObjects.Add(0, 'OLEDB;Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;'
                                       f'Location={query};Extended Properties=""', None, 1, ws.Range("$A$1"))
            lo.Name = query
            qt = lo.QueryTable
            qt.CommandType = 2
            qt.CommandText = f"SELECT * FROM [{query}]"
            qt.BackgroundQuery = False
            qt.Refresh(False)
            conn = qt.WorkbookConnection.OLEDBConnection
            conn.BackgroundQuery = True
            conn.RefreshOnFileOpen = True
            conn.RefreshPeriod = 1
        xl.CalculateFull()
        wb.Worksheets("How to use").Activate()
        wb.Save()
        wb.Close(False)
    finally:
        xl.Quit()


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--docs", type=Path, default=DOCS, help=argparse.SUPPRESS)
    ap.add_argument("--no-excel", action="store_true", help="skip the Power Query step")
    ap.add_argument("--migrate-from", type=Path, help="copy statuses from an old (pre-CSV) tracker xlsx")
    a = ap.parse_args(argv)

    docs = a.docs
    plan = build_plan(docs / "sera-sync-v3-blueprint.md")
    write_text_atomic(docs / "sera-sync-v3-plan.json", json.dumps(plan, indent=1, ensure_ascii=False) + "\n")

    status_csv, checks_csv = docs / "sera-sync-v3-status.csv", docs / "sera-sync-v3-checks.csv"
    if a.migrate_from:
        if status_csv.exists() or checks_csv.exists():
            raise SystemExit("Status CSVs already exist; refusing to overwrite them with a migration.")
        statuses, checks = migrate_from_xlsx(a.migrate_from)
        write_csv_atomic(status_csv, STATUS_FIELDS, statuses)
        write_csv_atomic(checks_csv, CHECK_FIELDS, checks)
        print(f"Migrated {sum(1 for s in statuses if s['Status'] != 'Not started')} started WPs and "
              f"{sum(1 for c in checks if c['Result'] != 'Not run')} recorded checks.")
    if not status_csv.exists():
        write_csv_atomic(status_csv, STATUS_FIELDS,
                         [{"WP": w["wp"], "Status": "Not started"} for w in plan["wps"]])
    if not checks_csv.exists():
        write_csv_atomic(checks_csv, CHECK_FIELDS,
                         [{"Check": c["n"], "Result": "Not run"} for c in plan["phase_checks"]])

    out = docs / "sera-sync-v3-agents.xlsx"
    if out.with_name("~$" + out.name).exists():
        raise SystemExit(f"{out.name} is open in Excel. Close it and run again.")
    build_viewer(plan, out)
    if not a.no_excel:
        attach_power_query(out, status_csv, checks_csv)
    print(f"Built {out.name}: {len(plan['wps'])} WPs, {len(plan['phase_checks'])} phase checks"
          + ("" if a.no_excel else ", Power Query attached (refresh on open + every minute)."))
    return 0


if __name__ == "__main__":
    sys.exit(main())
