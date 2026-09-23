"""
tools/sync_v3_tracker.py
------------------------
Updates docs/sera-sync-v3-agents.xlsx so agents working on the Sera Sync v3
blueprint never edit the spreadsheet by hand. Only the yellow input columns are
written; every formula is left alone.

  venv\\Scripts\\python tools\\sync_v3_tracker.py show P0-1
  venv\\Scripts\\python tools\\sync_v3_tracker.py set P0-1 --status "In progress" --model "Claude Haiku 4.5"
  venv\\Scripts\\python tools\\sync_v3_tracker.py set P0-1 --status "In review" --notes "debounce 60 s"
  venv\\Scripts\\python tools\\sync_v3_tracker.py set P0-3 --reviewed-by "Claude Opus 5.5"
  venv\\Scripts\\python tools\\sync_v3_tracker.py set P0-1 --status Done --commit 1a2b3c4
  venv\\Scripts\\python tools\\sync_v3_tracker.py checks
  venv\\Scripts\\python tools\\sync_v3_tracker.py check 7 --result Pass --notes "tested on reception PC"

Refuses (exit code 2) when an update breaks the blueprint's rules: unknown model,
model not allowed for the WP's tier, starting a WP whose dependencies are not Done,
or marking Done without a commit or without the required review.
"""

import argparse
import datetime
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
DEFAULT_FILE = REPO / "docs" / "sera-sync-v3-agents.xlsx"
STATUSES = ["Not started", "In progress", "In review", "Done", "Blocked"]
RESULTS = ["Not run", "Pass", "Fail"]
HEADER_ROW = 4


class TrackerError(Exception):
    pass


def _headers(ws, row=HEADER_ROW) -> dict:
    return {ws.cell(row=row, column=c).value: c for c in range(1, ws.max_column + 1)
            if ws.cell(row=row, column=c).value}


def _wp_rows(ws, cols) -> dict:
    rows = {}
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        wp = ws.cell(row=r, column=cols["WP"]).value
        if wp:
            rows[str(wp)] = r
    return rows


def _model_matrix(wb) -> dict:
    ms = wb["Models"]
    cols = _headers(ms)
    out = {}
    for r in range(HEADER_ROW + 1, ms.max_row + 1):
        name = ms.cell(row=r, column=cols["Model"]).value
        t1 = ms.cell(row=r, column=cols["T1"]).value
        if not name or t1 not in ("Yes", "No", "Caution"):
            break
        out[name] = {
            "T1": t1,
            "T2": ms.cell(row=r, column=cols["T2"]).value,
            "T3": ms.cell(row=r, column=cols["T3"]).value,
            "review": ms.cell(row=r, column=cols["Can review T3 work"]).value,
        }
    return out


def _check_writable(path: Path):
    lock = path.with_name("~$" + path.name)
    if lock.exists():
        raise TrackerError(f"{path.name} is open in Excel. Close it and run the command again.")


def _load(path: Path):
    from openpyxl import load_workbook
    if not path.exists():
        raise TrackerError(f"Tracker not found: {path}")
    return load_workbook(path)


def _save(wb, path: Path):
    _check_writable(path)
    try:
        wb.save(path)
    except PermissionError:
        raise TrackerError(f"{path.name} is locked (probably open in Excel). Close it and run the command again.")


def show(path: Path, wp: str | None):
    wb = _load(path)
    ws = wb["Work Packages"]
    cols = _headers(ws)
    rows = _wp_rows(ws, cols)
    wanted = [wp] if wp else list(rows)
    fields = ["WP", "Phase", "Tier", "What", "Depends on", "Review before merge", "Status",
              "Model used", "Reviewed by", "Commit", "Notes"]
    for w in wanted:
        if w not in rows:
            raise TrackerError(f"Unknown WP {w!r}. Valid: {', '.join(rows)}")
        r = rows[w]
        vals = {f: ws.cell(row=r, column=cols[f]).value for f in fields}
        unfinished = _unfinished_deps(ws, cols, rows, r)
        if wp:
            for f in fields:
                print(f"{f:>20}: {vals[f] if vals[f] is not None else ''}")
            print(f"{'Unfinished deps':>20}: {', '.join(unfinished) if unfinished else 'none'}")
        else:
            state = vals["Status"] or "Not started"
            if state == "Not started" and not unfinished:
                state += " (ready)"
            print(f"{w:<7} {vals['Tier']:<3} {state:<24} {vals['What']}")


def _unfinished_deps(ws, cols, rows, r) -> list:
    deps = ws.cell(row=r, column=cols["Depends on"]).value or ""
    out = []
    for d in [x.strip() for x in str(deps).split(",") if x.strip()]:
        if d not in rows:
            out.append(f"{d}(unknown)")
        elif ws.cell(row=rows[d], column=cols["Status"]).value != "Done":
            out.append(d)
    return out


def set_wp(path: Path, wp: str, status=None, model=None, reviewed_by=None, commit=None,
           notes=None, append_notes=False):
    wb = _load(path)
    ws = wb["Work Packages"]
    cols = _headers(ws)
    rows = _wp_rows(ws, cols)
    if wp not in rows:
        raise TrackerError(f"Unknown WP {wp!r}. Valid: {', '.join(rows)}")
    r = rows[wp]
    tier = ws.cell(row=r, column=cols["Tier"]).value
    models = _model_matrix(wb)
    messages = []

    if status is not None and status not in STATUSES:
        raise TrackerError(f"Status must be one of: {', '.join(STATUSES)}")

    for label, name in (("--model", model), ("--reviewed-by", reviewed_by)):
        if name is not None and name not in models:
            raise TrackerError(f"{label}: unknown model {name!r}. Use a name from the Models sheet: "
                               f"{', '.join(models)}")

    if model is not None:
        allowed = models[model][tier]
        if allowed == "No":
            raise TrackerError(f"{model} is not allowed for {tier} work ({wp}). See the Models sheet.")
        if allowed == "Caution":
            messages.append(f"NOTE: {model} on {tier} needs a Fable 5.1 / Opus 5.5 review before merge.")

    if reviewed_by is not None and models[reviewed_by]["review"] != "Yes":
        raise TrackerError(f"{reviewed_by} can't be the reviewer. Reviews need Claude Fable 5.1 or Opus 5.5.")

    if status in ("In progress", "In review", "Done"):
        unfinished = _unfinished_deps(ws, cols, rows, r)
        if unfinished:
            raise TrackerError(f"{wp} depends on work that isn't Done yet: {', '.join(unfinished)}")

    if status == "Done":
        final_commit = commit if commit is not None else ws.cell(row=r, column=cols["Commit"]).value
        if not final_commit:
            raise TrackerError("Can't mark Done without a commit hash (--commit).")
        used = model if model is not None else ws.cell(row=r, column=cols["Model used"]).value
        if not used:
            raise TrackerError("Can't mark Done without the model that did the work (--model).")
        needs_review = str(ws.cell(row=r, column=cols["Review before merge"]).value or "").startswith("Yes")
        if used in models and models[used][tier] == "Caution":
            needs_review = True
        reviewer = reviewed_by if reviewed_by is not None else ws.cell(row=r, column=cols["Reviewed by"]).value
        if needs_review and not reviewer:
            raise TrackerError(f"{wp} needs a Fable 5.1 / Opus 5.5 review before Done (--reviewed-by).")

    updates = {"Status": status, "Model used": model, "Reviewed by": reviewed_by, "Commit": commit}
    for field, value in updates.items():
        if value is not None:
            ws.cell(row=r, column=cols[field]).value = value
    if notes is not None:
        stamp = datetime.date.today().isoformat()
        cell = ws.cell(row=r, column=cols["Notes"])
        new = f"{stamp}: {notes}"
        cell.value = f"{cell.value}\n{new}" if (append_notes and cell.value) else new

    _save(wb, path)
    for m in messages:
        print(m)
    print(f"Updated {wp}.")


def list_checks(path: Path):
    wb = _load(path)
    pc = wb["Phase Checks"]
    cols = _headers(pc, row=3)
    for r in range(4, pc.max_row + 1):
        chk = pc.cell(row=r, column=cols["Check"]).value
        if chk:
            n = r - 3
            print(f"{n:>3}  phase {pc.cell(row=r, column=cols['Phase']).value}  "
                  f"[{pc.cell(row=r, column=cols['Result']).value}]  {chk}")


def set_check(path: Path, number: int, result: str, notes=None):
    if result not in RESULTS:
        raise TrackerError(f"--result must be one of: {', '.join(RESULTS)}")
    wb = _load(path)
    pc = wb["Phase Checks"]
    cols = _headers(pc, row=3)
    r = number + 3
    if number < 1 or not pc.cell(row=r, column=cols["Check"]).value:
        raise TrackerError(f"No phase check number {number}. Run 'checks' to list them.")
    pc.cell(row=r, column=cols["Result"]).value = result
    pc.cell(row=r, column=cols["Date"]).value = datetime.date.today().isoformat()
    if notes is not None:
        pc.cell(row=r, column=cols["Notes"]).value = notes
    _save(wb, path)
    print(f"Updated check {number}.")


def main(argv=None) -> int:
    # The sheets contain characters like "→" that the default Windows console code page can't print
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description="Update the Sera Sync v3 agent tracker spreadsheet.")
    ap.add_argument("--file", type=Path, default=DEFAULT_FILE, help=argparse.SUPPRESS)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_show = sub.add_parser("show", help="show all WPs, or one WP in detail")
    p_show.add_argument("wp", nargs="?")

    p_set = sub.add_parser("set", help="update a WP's yellow columns")
    p_set.add_argument("wp")
    p_set.add_argument("--status", choices=STATUSES)
    p_set.add_argument("--model")
    p_set.add_argument("--reviewed-by")
    p_set.add_argument("--commit")
    p_set.add_argument("--notes")
    p_set.add_argument("--append-notes", action="store_true", help="add to existing notes instead of replacing")

    sub.add_parser("checks", help="list the hands-on phase checks")
    p_chk = sub.add_parser("check", help="record the result of a phase check")
    p_chk.add_argument("number", type=int)
    p_chk.add_argument("--result", required=True, choices=RESULTS)
    p_chk.add_argument("--notes")

    a = ap.parse_args(argv)
    try:
        if a.cmd == "show":
            show(a.file, a.wp)
        elif a.cmd == "set":
            set_wp(a.file, a.wp, a.status, a.model, a.reviewed_by, a.commit, a.notes, a.append_notes)
        elif a.cmd == "checks":
            list_checks(a.file)
        elif a.cmd == "check":
            set_check(a.file, a.number, a.result, a.notes)
    except TrackerError as e:
        print(f"REFUSED: {e}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
