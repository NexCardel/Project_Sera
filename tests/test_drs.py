"""tests/test_drs.py - Unit and integration tests for Deadline Reminder System (DRS).

Verifies per-form column customization, full MCL roster population,
comprehensive tracker dump fields, 'extract from tracker dump only' (LTT mode),
and multi-sheet Excel generation.
"""
from datetime import date
from pathlib import Path

from core.drs import config, engine, export
from core.drs.config import SCOPE_ALL_MCL, SCOPE_SERVICE_ONLY


class DummyDB:
    def __init__(self):
        self.app_dir = Path(".")
        self._settings = {}
        self._services = [
            {"id": 1, "name": "GST", "login_page_link": "https://services.gst.gov.in"},
            {"id": 2, "name": "Income Tax", "login_page_link": "https://eportal.incometax.gov.in"},
        ]
        self._mcl = [
            {"id": 1, "label": "Client Name", "field_type": "text", "is_identity": 0},
            {"id": 2, "label": "PAN", "field_type": "text", "is_identity": 1},
            {"id": 3, "label": "GSTIN", "field_type": "text", "is_identity": 0},
            {"id": 4, "label": "Password", "field_type": "password", "is_identity": 0},
            {"id": 5, "label": "File No", "field_type": "text", "is_identity": 0},
            {"id": 6, "label": "Category", "field_type": "text", "is_identity": 0},
        ]
        self._clients = [
            {
                "id": 101,
                "values": {1: "Alpha Traders", 2: "ABCDE1234F", 3: "27ABCDE1234F1Z5", 5: "F-01", 6: "Retail"},
                "service_ids": [1],  # GST only
            },
            {
                "id": 102,
                "values": {1: "Beta Services", 2: "BCDEF2345G", 5: "F-02", 6: "Corporate"},
                "service_ids": [2],  # Income Tax only
                "notes": "Beta client note from Master DB",
            },
            {
                "id": 103,
                "values": {1: "Gamma Industries", 2: "CDEFG3456H", 3: "27CDEFG3456H1Z2", 5: "F-03", 6: "Manufacturing"},
                "service_ids": [1, 2],  # Both
            },
        ]
        self._tracker_dumps = []
        self._containers = [
            {
                "identity_key": "ABCDE1234F",
                "client_id": 101,
                "pan": "ABCDE1234F",
                "gstin": "27ABCDE1234F1Z5",
                "company_name": "Alpha Traders",
                "filing_history": [
                    {
                        "portal": "GST (GSTR-3B)",
                        "period_label": "September (FY 2026-27)",
                        "status": "Filed",
                        "arn": "AA2709260000001",
                        "created_at": "2026-10-20T10:00:00Z",
                        "captured_by": "PC-01",
                        "capture_method": "DOM_Tracker",
                        "original_author": "PC-01",
                        "notes": "Alpha note stored in JSON",
                    },
                    {
                        "portal": "GST (GSTR-3B)",
                        "period_label": "September (FY 2026-27)",
                        "status": "Draft",
                        "arn": "N/A",
                        "created_at": "2026-10-21T10:00:00Z",
                        "captured_by": "PC-01",
                        "capture_method": "DOM_Tracker",
                        "original_author": "PC-01",
                        "notes": "Alpha note stored in JSON",
                    },
                ],
            }
        ]

    def get_setting(self, key, default=None):
        return self._settings.get(key, default)

    def set_setting(self, key, val):
        self._settings[key] = val

    def get_services(self):
        return self._services

    def get_service_for_portal(self, portal_name):
        clean = str(portal_name).lower()
        if "gst" in clean:
            return self._services[0]
        if "income" in clean or "itr" in clean:
            return self._services[1]
        return None

    def get_mcl_columns(self):
        return self._mcl

    def search_clients(self, query="", include_archived=False):
        return self._clients

    def get_srpf_containers(self, limit=1_000_000, slim=True):
        return self._containers

    def get_tracker_dumps(self, client_id=None, limit=200, search_query=None, slim=False, dump_id=None):
        return self._tracker_dumps


def test_drs_default_columns():
    db = DummyDB()
    cols = config.default_drs_columns(db)
    col_labels = [c["label"] for c in cols]

    # Password column must NEVER be exposed
    assert "Password" not in col_labels

    # PAN, GSTIN, Name, Fixed columns and Tracker fields present
    assert "PAN" in col_labels
    assert "Client Name" in col_labels
    assert "Latest period" in col_labels
    assert "Status" in col_labels
    assert "ARN" in col_labels

    # Full MCL roster fields present (File No, Category)
    assert "File No" in col_labels
    assert "Category" in col_labels

    # Full suite of Tracker Dump fields present
    assert "Captured on" in col_labels
    assert "Captured by" in col_labels
    assert "Capture method" in col_labels
    assert "Original author" in col_labels
    assert "Portal Status" in col_labels
    assert "Filing Period" in col_labels
    assert "Portal ID" in col_labels
    assert "Dataset Key" in col_labels
    assert "Notes" in col_labels

    # Non-GST form should have GSTIN disabled by default
    itr_cols = config.default_drs_columns_for_form(db, "ITR-1", is_gst=False)
    gstin_col = next(c for c in itr_cols if c["label"] == "GSTIN")
    assert gstin_col["enabled"] is False


def test_drs_config_load_save():
    db = DummyDB()
    cfg = config.load_drs_config(db, detected_forms=["GSTR-3B", "ITR-1"])
    assert "forms" in cfg
    assert "GSTR-3B" in cfg["forms"]
    assert "ITR-1" in cfg["forms"]
    assert cfg["client_roster_scope"] == SCOPE_ALL_MCL

    # Customize GSTR-3B columns and set tracker_only
    cfg["forms"]["GSTR-3B"]["tracker_only"] = True
    cfg["forms"]["GSTR-3B"]["columns"][0]["enabled"] = False
    config.save_drs_config(db, cfg)

    reloaded = config.load_drs_config(db, detected_forms=["GSTR-3B", "ITR-1"])
    assert reloaded["forms"]["GSTR-3B"]["tracker_only"] is True
    assert reloaded["forms"]["GSTR-3B"]["columns"][0]["enabled"] is False


def test_form_detection_and_service_linking():
    db = DummyDB()
    detected = engine.detect_forms_and_services(db)
    assert len(detected) >= 1

    form_names = [f["form"] for f in detected]
    assert "GSTR-3B" in form_names

    gstr3b = next(f for f in detected if f["form"] == "GSTR-3B")
    assert gstr3b["portal"] == "GST"
    assert gstr3b["service"]["name"] == "GST"
    assert gstr3b["is_gst"] is True


def test_build_drs_dataset_whole_mcl_roster():
    """Default mode (SCOPE_ALL_MCL) brings the entire client roster to the form sheet."""
    db = DummyDB()
    dataset = engine.build_drs_dataset(db)

    assert "GSTR-3B" in dataset
    rows = dataset["GSTR-3B"]

    # All three clients from MCL appear on GSTR-3B
    client_names = [r.get("Client Name") for r in rows]
    assert "Alpha Traders" in client_names
    assert "Beta Services" in client_names
    assert "Gamma Industries" in client_names

    # Alpha Traders has capture
    alpha_row = next(r for r in rows if r.get("Client Name") == "Alpha Traders")
    assert alpha_row["Latest period"] == "September (FY 2026-27)"
    assert alpha_row["Status"] == "Submitted & Verified"
    assert alpha_row["ARN"] == "AA2709260000001"

    # Beta and Gamma have NO captures yet:
    beta_row = next(r for r in rows if r.get("Client Name") == "Beta Services")
    assert beta_row["Latest period"] == "no captures yet"
    assert beta_row["Status"] == "no captures yet"

    gamma_row = next(r for r in rows if r.get("Client Name") == "Gamma Industries")
    assert gamma_row["Latest period"] == "no captures yet"
    assert gamma_row["Status"] == "no captures yet"


def test_build_drs_dataset_service_only_scope():
    """Service-enrolled scope only populates clients attached to that service."""
    db = DummyDB()
    cfg = config.load_drs_config(db)
    cfg["client_roster_scope"] = SCOPE_SERVICE_ONLY
    config.save_drs_config(db, cfg)

    dataset = engine.build_drs_dataset(db, config=cfg)
    rows = dataset["GSTR-3B"]
    client_names = [r.get("Client Name") for r in rows]

    assert "Alpha Traders" in client_names
    assert "Gamma Industries" in client_names
    # Beta Services (Income Tax only) must NOT appear in GSTR-3B under service_only scope
    assert "Beta Services" not in client_names


def test_build_drs_dataset_tracker_only_ltt_mode():
    """'Extract from tracker dump only' (LTT mode) omits uncaptured clients entirely."""
    db = DummyDB()
    cfg = config.load_drs_config(db, detected_forms=["GSTR-3B"])
    cfg["forms"]["GSTR-3B"]["tracker_only"] = True
    config.save_drs_config(db, cfg)

    dataset = engine.build_drs_dataset(db, config=cfg)
    rows = dataset["GSTR-3B"]
    client_names = [r.get("Client Name") for r in rows]

    # Only Alpha Traders has a capture for GSTR-3B; uncaptured clients are omitted
    assert "Alpha Traders" in client_names
    assert "Gamma Industries" not in client_names
    assert "Beta Services" not in client_names


def test_per_form_column_mapping():
    """Different filing forms can have distinct column configurations."""
    db = DummyDB()
    cfg = config.load_drs_config(db, detected_forms=["GSTR-3B"])
    # Enable Capture method, Original author, File No, and Notes on GSTR-3B
    for c in cfg["forms"]["GSTR-3B"]["columns"]:
        if c["label"] in ("Capture method", "Original author", "File No", "Notes"):
            c["enabled"] = True
    config.save_drs_config(db, cfg)

    dataset = engine.build_drs_dataset(db, config=cfg)
    alpha = next(r for r in dataset["GSTR-3B"] if r.get("Client Name") == "Alpha Traders")
    assert alpha.get("Capture method") == "DOM_Tracker"
    assert alpha.get("Original author") == "PC-01"
    assert alpha.get("File No") == "F-01"
    assert alpha.get("Notes") == "Alpha note stored in JSON"

    # Beta Services (uncaptured, whole MCL roster) receives its Master DB client notes
    beta = next(r for r in dataset["GSTR-3B"] if r.get("Client Name") == "Beta Services")
    assert beta.get("Notes") == "Beta client note from Master DB"


def test_tracker_dump_raw_payload_json_notes_flow_to_drs():
    """Notes stored inside raw_payload_json connect directly to DRS rows."""
    import json
    db = DummyDB()
    db._containers = []
    db._tracker_dumps = [{
        "client_id": 103,
        "unassigned_identity": "CDEFG3456H",
        "portal": "GST (GSTR-3B)",
        "period_label": "September (FY 2026-27)",
        "arn_number": "AA2709260000099",
        "status": "Filed",
        "created_at": "2026-10-20T10:00:00Z",
        "captured_by": "PC-02",
        "raw_payload_json": json.dumps({
            "status": "Filed",
            "arn": "AA2709260000099",
            "notes": "Payment verified via challan 12345",
        }),
    }]

    cfg = config.load_drs_config(db, detected_forms=["GSTR-3B"])
    for c in cfg["forms"]["GSTR-3B"]["columns"]:
        if c["label"] == "Notes":
            c["enabled"] = True
    config.save_drs_config(db, cfg)

    dataset = engine.build_drs_dataset(db, config=cfg)
    gamma = next(r for r in dataset["GSTR-3B"] if r.get("Client Name") == "Gamma Industries")
    assert gamma["Notes"] == "Payment verified via challan 12345"


def test_save_tracker_dump_media_updates_raw_payload_json(tmp_path):
    """Saving tracker media injects 'notes' into the raw_payload_json for DRS and sync."""
    import json
    import sqlite3
    from sera_db.tracker_dump import TrackerDumpMixin

    db_file = str(tmp_path / "rawPayload.db")

    class TestDB(TrackerDumpMixin):
        def _connect_raw(self):
            return sqlite3.connect(db_file)

    conn = sqlite3.connect(db_file)
    conn.execute("CREATE TABLE tracker_dump (id INTEGER PRIMARY KEY, notes TEXT, screenshot_path TEXT, raw_payload_json TEXT)")
    conn.execute("INSERT INTO tracker_dump (id, notes, screenshot_path, raw_payload_json) VALUES (1, '', '', '{\"status\": \"Filed\"}')")
    conn.commit()
    conn.close()

    db = TestDB()
    ok = db.save_tracker_dump_media(1, "Challan paid and confirmed", "media/snap.png")
    assert ok is True

    conn = sqlite3.connect(db_file)
    row = conn.execute("SELECT notes, screenshot_path, raw_payload_json FROM tracker_dump WHERE id = 1").fetchone()
    conn.close()

    assert row[0] == "Challan paid and confirmed"
    assert row[1] == "media/snap.png"
    parsed = json.loads(row[2])
    assert parsed.get("notes") == "Challan paid and confirmed"


def test_export_and_workbook_creation(tmp_path):
    db = DummyDB()
    db.app_dir = tmp_path

    # Export DRS (excel=False exercises static fallback deterministically)
    xlsx = export.export_drs(db, app_dir=tmp_path, today=date(2026, 10, 22),
                             force_rebuild=True, excel=False)
    assert xlsx.exists()
    assert xlsx.name == "DRS_2026-10.xlsx"

    # Both companion CSV feeds must be written next to the workbook
    data_csv, over_csv = export.drs_csv_paths(xlsx)
    assert data_csv.exists()
    assert over_csv.exists()

    # Content verification
    csv_text = data_csv.read_text(encoding="utf-8")
    assert "Alpha Traders" in csv_text
    assert "Gamma Industries" in csv_text
    assert "no captures yet" in csv_text

    # Native filter arrows must be present in static fallback
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["GSTR-3B"]
    assert ws.auto_filter.ref.startswith("A3:")
    assert wb["Overview"].auto_filter.ref is not None

    # LTT-style layout checks
    assert wb.properties.title == export.FORMAT
    assert wb["Overview"]["A1"].value == "DRS — October 2026"
    assert wb["Overview"].column_dimensions["I"].hidden is True
    assert ws.freeze_panes == "A4"
    assert ws.print_title_rows in ("3:3", "$3:$3")
    assert ws["A1"].font.size == 16
    assert ws["A3"].fill.start_color.rgb == export.HEAD_FILL

    from openpyxl.utils import get_column_letter
    labels = [c["label"] for c in export.get_active_columns_for_form(db, "GSTR-3B")]
    col_of = lambda label: get_column_letter(labels.index(label) + 1)  # noqa: E731
    name_cell = ws[f"{col_of('Client Name')}4"]
    assert name_cell.font.name == "Calibri"
    assert ws[f"{col_of('PAN')}4"].font.name == "Consolas"
    assert name_cell.fill.start_color.rgb == export.BAND_FILL
    status_cell = ws[f"{col_of('Status')}4"]
    assert status_cell.value == "Submitted & Verified"
    assert status_cell.fill.start_color.rgb == export.STATUS_FILL["Submitted & Verified"]


def test_live_workbook_refreshes_when_it_is_open(tmp_path, monkeypatch):
    db = DummyDB()
    db.app_dir = tmp_path
    xlsx = export.export_drs(db, app_dir=tmp_path, today=date(2026, 10, 22), excel=False)

    refreshed = []
    monkeypatch.setattr(export, "_has_live_tables", lambda _p: True)
    monkeypatch.setattr(export, "_open_in_excel", lambda _p: True)
    monkeypatch.setattr(export, "refresh_open_drs_workbook", lambda p: refreshed.append(p) or True)

    assert export.export_drs(db, app_dir=tmp_path, today=date(2026, 10, 22)) == xlsx
    assert refreshed == [xlsx]


def test_tracker_dump_fallback_covers_unaggregated_capture():
    db = DummyDB()
    db._containers = []
    db._tracker_dumps = [{
        "client_id": 103,
        "unassigned_identity": "CDEFG3456H",
        "portal": "GST (GSTR-3B)",
        "period_label": "September (FY 2026-27)",
        "arn_number": "AA2709260000099",
        "status": "Filed",
        "created_at": "2026-10-20T10:00:00Z",
        "captured_by": "PC-02",
    }]

    dataset = engine.build_drs_dataset(db)
    gamma = next(r for r in dataset["GSTR-3B"] if r.get("Client Name") == "Gamma Industries")
    assert gamma["Latest period"] == "September (FY 2026-27)"
    assert gamma["Status"] == "Submitted & Verified"
    assert gamma["ARN"] == "AA2709260000099"


def test_existing_workbook_rebuilds_when_a_new_form_appears(tmp_path):
    db = DummyDB()
    db.app_dir = tmp_path
    first = export.export_drs(db, app_dir=tmp_path, today=date(2026, 10, 22), excel=False)

    db._containers.append({
        "identity_key": "BCDEF2345G",
        "client_id": 102,
        "pan": "BCDEF2345G",
        "company_name": "Beta Services",
        "filing_history": [{
            "portal": "Income Tax (ITR-1)",
            "period_label": "AY 2026-27",
            "status": "Filed",
            "arn": "AA2709260000088",
            "created_at": "2026-10-21T10:00:00Z",
            "captured_by": "PC-01",
        }],
    })

    export.export_drs(db, app_dir=tmp_path, today=date(2026, 10, 22), excel=False)
    from openpyxl import load_workbook
    wb = load_workbook(first)
    assert "GSTR-3B" in wb.sheetnames
    assert "ITR-1" in wb.sheetnames


def test_scheduled_export_debounce():
    db = DummyDB()
    db.app_dir = Path(".")
    t1 = export.schedule_drs_export(db, delay=0.1)
    t2 = export.schedule_drs_export(db, delay=0.1)
    assert t2 is not None
    assert t1 is not t2
    t2.cancel()
