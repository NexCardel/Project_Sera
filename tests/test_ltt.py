from datetime import date, datetime
from pathlib import Path

from core.ltt import feed, rules
from core.ltt.engine import LTTEngine, parse_date, parse_gap
from core.sgt.sgt_toolbox import SUBMIT_LEVELS

TODAY = date(2026, 10, 4)


def _c(pan, name, *hist, company=""):
    return {"pan": pan, "company_name": company, "proprietor_name": name, "gstin": "",
            "filing_history": [dict(zip(("portal", "period_label", "status", "arn", "created_at"), h)) for h in hist]}


def test_levels_match_sgt_ladder():
    assert feed.LEVELS == tuple(SUBMIT_LEVELS)


def test_engine_label_is_trackers_wording():
    e = LTTEngine(parse_date("20-10-2026"), parse_gap("monthly"), 1, 4)
    assert e.locate(TODAY)["current"].filing == "September (FY 2026-27)"
    assert e.locate(TODAY)["previous"][0].filing == "August (FY 2026-27)"


def test_fy_end_lag():
    r = rules.FormRule("ITR", "Yearly", "2026-07-31", "fy_end")
    assert r.engine().locate(TODAY)["current"].filing == "March (FY 2026-27)"   # due Jul 2027
    assert r.engine().locate(date(2026, 6, 1))["current"].filing == "March (FY 2025-26)"
    assert rules.FormRule("GSTR-9", "Yearly", "2026-12-31", "fy_end").lag() == 9


def test_ladder():
    assert feed.ladder("Not e-verified") == "Submitted (Not Verified)"
    assert feed.ladder("Filed") == "Submitted & Verified"
    assert feed.ladder("Form Selected") == "Draft"
    assert feed.ladder("", "N/A") == "Not Submitted"
    assert feed.ladder("Visited", "123456789012345") == "Submitted & Verified"
    assert feed.ladder("Option Expired") == feed.NOT_APPLICABLE


def test_parse_period():
    assert feed.parse_period("June (FY 2026-27)") == ({6}, 2026)
    assert feed.parse_period("Jul-Sep (FY 2026-27)") == ({7, 8, 9}, 2026)
    assert feed.parse_period("AY 2026-27") == (set(), 2025)
    assert feed.parse_period("Q2 FY 2026-27") == ({7, 8, 9}, 2026)
    assert feed.parse_period("") == (set(), None)


def test_form_matching():
    assert rules.form_matches("GSTR-3B", "GSTR3B")
    assert rules.form_matches("ITR", "ITR-4")
    assert not rules.form_matches("GSTR-1", "GSTR-1A")
    assert not rules.form_matches("GSTR-1", "GSTR-3B")


def test_build_rows():
    rs = [rules.FormRule("GSTR-3B", "Monthly", "2026-10-20"),
          rules.FormRule("ITR", "Yearly", "2026-07-31", "fy_end")]
    cs = [
        _c("ABCDE1234F", "Asha", ("GST (GSTR-3B)", "September (FY 2026-27)", "Not e-verified", "A1", "2026-10-03T08:30:00+00:00"),
           ("GST (GSTR-3B)", "August (FY 2026-27)", "Filed", "A0"), company="Asha Traders"),
        _c("PQRST5678K", "Ravi", ("GST (GSTR-3B)", "August (FY 2026-27)", "Filed", "B0"),
           ("Income Tax (ITR-4)", "AY 2026-27", "Draft", "N/A")),
    ]
    rows = {(r["PAN"], r["Return form"]): r for r in feed.build_rows(cs, rs, TODAY)}
    a = rows[("ABCDE1234F", "GSTR-3B")]
    assert (a["Name"], a["Current period"], a["Submit status"]) == (
        "Asha Traders", "September (FY 2026-27)", "Submitted (Not Verified)")
    assert a["Capture date & time"].startswith("2026-10-0")                   # local time of that capture
    assert rows[("PQRST5678K", "GSTR-3B")]["Capture date & time"] == ""
    assert rows[("PQRST5678K", "GSTR-3B")]["Submit status"] == "Not Submitted"   # only August filed
    # ITR current period is March (FY 2026-27) -> AY 2027-28, so AY 2026-27 does not count
    assert rows[("PQRST5678K", "ITR")]["Submit status"] == "Not Submitted"
    assert ("ABCDE1234F", "ITR") not in rows


def test_include_exclude():
    r = rules.FormRule("GSTR-3B", "Monthly", "2026-10-20", include=["NEWCL1234A"], exclude=["ABCDE1234F"])
    cs = [_c("ABCDE1234F", "Asha", ("GST (GSTR-3B)", "August (FY 2026-27)", "Filed", "")),
          _c("NEWCL1234A", "New")]
    pans = [x["PAN"] for x in feed.build_rows(cs, [r], TODAY)]
    assert pans == ["NEWCL1234A"]


def test_rules_roundtrip(tmp_path):
    assert [r.form for r in rules.load_rules(tmp_path)][0] == "GSTR-1"        # defaults
    rs = [rules.FormRule("GSTR-3B", "Monthly", "2026-10-20", include=["abcde1234f"])]
    rules.save_rules(tmp_path, rs)
    assert rules.load_rules(tmp_path)[0].include == ["ABCDE1234F"]
    assert rules.validate([rules.FormRule("X", "Monthly", "bad")]) != ""
    feed.write_csv(tmp_path / "f.csv", [])
    assert (tmp_path / "f.csv").read_text(encoding="utf-8-sig").strip() == ",".join(feed.FIELDS)


# ---- engine: grace, extensions, holidays, month-end, covers ----

def _eng(deadline, gap, **kw):
    return LTTEngine(parse_date(deadline), parse_gap(gap), **kw)


def test_grace_keeps_overdue_period_current():
    e = _eng("2026-08-31", "yearly", covers="fy_end", grace_days=122)
    assert e.locate(TODAY)["current"].filing == "March (FY 2025-26)"
    assert e.locate(date(2026, 12, 31))["current"].filing == "March (FY 2025-26)"
    assert e.locate(date(2027, 1, 1))["current"].filing == "March (FY 2026-27)"
    assert _eng("2026-08-31", "yearly", covers="fy_end").locate(TODAY)["current"].filing == "March (FY 2026-27)"


def test_extension_moves_deadline_not_filing_month():
    e = _eng("2026-10-20", "monthly", overrides={date(2026, 10, 20): date(2026, 11, 3)})
    cur = e.locate(date(2026, 10, 25))["current"]
    assert cur.end == date(2026, 11, 3) and cur.filing == "September (FY 2026-27)"
    assert e.locate(date(2026, 11, 4))["current"].filing == "October (FY 2026-27)"
    assert e.locate(date(2026, 11, 4))["current"].start == date(2026, 11, 4)


def test_itr_extended_into_next_month_still_files_march():
    e = _eng("2026-07-31", "yearly", covers="fy_end", overrides={date(2026, 7, 31): date(2026, 9, 15)})
    assert e.locate(date(2026, 8, 20))["current"].filing == "March (FY 2025-26)"


def test_holidays_and_sundays_roll_forward():
    e = _eng("2026-10-20", "monthly", holidays=[date(2026, 11, 20)], roll_sundays=True)
    assert e.deadline(1) == date(2026, 11, 21)       # Fri holiday -> Sat
    assert e.deadline(3) == date(2027, 1, 20)        # Wed untouched
    assert e.deadline(2) == date(2026, 12, 21)       # Sun 20 Dec -> Mon
    assert _eng("2026-10-20", "monthly").deadline(2) == date(2026, 12, 20)   # off by default


def test_month_end_follows_month_length():
    e = _eng("2026-04-30", "monthly", month_end=True)
    assert [e.deadline(k) for k in (1, 2, 10)] == [date(2026, 5, 31), date(2026, 6, 30), date(2027, 2, 28)]
    assert _eng("2026-04-30", "monthly").deadline(1) == date(2026, 5, 30)    # default unchanged


def test_covers_calendar_year_and_validation():
    e = _eng("2026-03-31", "yearly", covers="fy_end", fy_start=1)
    assert e.locate(date(2026, 3, 1))["current"].filing == "December (FY 2025)"
    assert _eng("2026-10-20", "monthly", covers="same_month").locate(TODAY)["current"].filing == "October (FY 2026-27)"
    for bad in ({"covers": "nope"}, {"grace_days": -1}):
        try:
            _eng("2026-10-20", "monthly", **bad)
            assert False
        except ValueError:
            pass


# ---- monthly workbook: latest period per client per form ----
from core.ltt import monthly   # noqa: E402


def test_latest_period_only_and_split_by_form():
    cs = [_c("ABCDE1234F", "Asha",
             ("GST (GSTR-3B)", "June (FY 2026-27)", "Filed", "A1", "2026-07-03T08:30:00+00:00"),
             ("GST (GSTR-3B)", "September (FY 2026-27)", "Not e-verified", "A4", "2026-10-03T08:30:00+00:00"),
             ("GST (GSTR3B)", "March (FY 2025-26)", "Filed", "A9", "2026-10-05T08:30:00+00:00"),      # late capture, old period
             ("GST (GSTR-1)", "August (FY 2026-27)", "Draft", "N/A", "2026-09-10T08:30:00+00:00"),
             ("Income Tax (ITR-4)", "AY 2025-26", "Verified", "X", "2025-08-01T00:00:00+00:00"),
             ("Income Tax (ITR-4)", "AY 2026-27", "Draft", "N/A", "2026-08-01T00:00:00+00:00"),
             company="Asha Traders")]
    d = monthly.latest_rows(cs)
    assert list(d) == ["GSTR-1", "GSTR-3B", "ITR-4"]
    assert len(d["GSTR-3B"]) == 1
    r = d["GSTR-3B"][0]
    assert (r["Latest period"], r["Status"], r["ARN"]) == ("September (FY 2026-27)", "Submitted (Not Verified)", "A4")
    assert d["GSTR-1"][0]["Status"] == "Draft" and d["GSTR-1"][0]["ARN"] == ""
    assert (d["ITR-4"][0]["Latest period"], d["ITR-4"][0]["Status"]) == ("AY 2026-27", "Draft")


def test_month_cutoff_freezes_view():
    cs = [_c("ABCDE1234F", "Asha",
             ("GST (GSTR-3B)", "August (FY 2026-27)", "Filed", "A1", "2026-09-20T08:30:00+00:00"),
             ("GST (GSTR-3B)", "September (FY 2026-27)", "Filed", "A2", "2026-10-03T08:30:00+00:00"))]
    assert monthly.latest_rows(cs, date(2026, 9, 1))["GSTR-3B"][0]["Latest period"] == "August (FY 2026-27)"
    assert monthly.latest_rows(cs, date(2026, 10, 1))["GSTR-3B"][0]["Latest period"] == "September (FY 2026-27)"
    assert monthly.latest_rows(cs, date(2026, 8, 1)) == {}


def test_period_end_ordering():
    assert monthly.period_end("Jul-Sep (FY 2026-27)") == date(2026, 9, 1)
    assert monthly.period_end("January (FY 2026-27)") == date(2027, 1, 1)
    assert monthly.period_end("AY 2026-27", True) == date(2026, 3, 1)
    assert monthly.period_end("September") is None


def _sample():
    cs = [_c(f"ABCDE123{i}F", f"Client {i}", ("GST (GSTR-3B)", "September (FY 2026-27)", "Filed", "A")) for i in range(3)]
    cs.append(_c("PQRST5678K", "Ravi", ("Income Tax (ITR-4)", "AY 2026-27", "Draft", "N/A")))
    return monthly.latest_rows(cs, date(2026, 10, 1))


def test_csvs_hold_every_form_and_a_summary_row_per_form(tmp_path):
    import csv
    x = tmp_path / "LTT" / "LTT_2026-10.xlsx"
    monthly.write_csvs(x, _sample())
    data_csv, over_csv = monthly.csv_paths(x)
    rows = list(csv.DictReader(open(data_csv, encoding="utf-8-sig")))
    assert [r["Form"] for r in rows] == ["GSTR-3B"] * 3 + ["ITR-4"]
    over = {r["Form"]: r for r in csv.DictReader(open(over_csv, encoding="utf-8-sig"))}
    assert over["GSTR-3B"]["Clients"] == "3" and over["GSTR-3B"]["Submitted & Verified"] == "3"
    assert over["ITR-4"]["Latest period"] == "AY 2026-27" and over["ITR-4"]["Draft"] == "1"


def test_shell_has_a_sheet_per_form_and_is_rebuilt_when_a_form_appears(tmp_path):
    x = tmp_path / "LTT_2026-10.xlsx"
    forms = list(_sample())
    monthly.build_shell(x, date(2026, 10, 1), forms)
    assert monthly.workbook_matches(x, forms)
    assert not monthly.workbook_matches(x, forms + ["GSTR-1"])
    from openpyxl import load_workbook
    wb = load_workbook(x)
    assert wb.sheetnames == ["Overview", "GSTR-3B", "ITR-4"]
    assert wb["Overview"]["A1"].value == "LTT — October 2026"
    assert wb["GSTR-3B"].max_row == 2          # nothing pre-filled: the tables arrive with the queries


def test_old_layout_is_rebuilt_and_a_stale_lock_file_is_ignored(tmp_path):
    from openpyxl import Workbook, load_workbook
    x = tmp_path / "LTT_2026-10.xlsx"
    old = Workbook()
    old.active.title = "Overview"
    old.save(x)                                                   # no layout stamp = an older build
    (tmp_path / "~$LTT_2026-10.xlsx").write_text("")              # leftover owner file, nobody has it open
    assert not monthly.workbook_matches(x, [])
    monthly.ensure_workbook(x, date(2026, 10, 1), ["GSTR-3B"], excel=False)
    assert monthly.workbook_matches(x, ["GSTR-3B"])
    assert load_workbook(x).properties.title == monthly.FORMAT


# ---- client names: registered name > captured trade name > proprietor name ----

def test_clean_name_strips_labels_only():
    assert feed.clean_name("Trade Name -") == ""
    assert feed.clean_name("Legal Name :") == ""
    assert feed.clean_name("-") == ""
    assert feed.clean_name("Trade Name - ABC Traders") == "ABC Traders"
    assert feed.clean_name("Namrata Traders") == "Namrata Traders"


def test_name_priority_registered_then_trade_then_proprietor():
    both = {"registered_name": "Saved Co", "company_name": "Captured Trade", "proprietor_name": "Owner"}
    assert feed.name_of(both) == ("Saved Co", 0)
    assert feed.name_of({"company_name": "Captured Trade", "proprietor_name": "Owner"}) == ("Captured Trade", 1)
    assert feed.name_of({"company_name": "Trade Name -", "proprietor_name": "Owner"}) == ("Owner", 2)
    assert feed.name_of({"registered_name": "", "company_name": "", "proprietor_name": ""}) == ("", 9)


def test_registered_name_wins_when_containers_merge():
    a = {"pan": "ABCDE1234F", "company_name": "Trade Name -", "proprietor_name": "Asha", "filing_history": []}
    b = {"pan": "ABCDE1234F", "company_name": "", "proprietor_name": "", "registered_name": "Asha Traders Pvt Ltd",
         "filing_history": []}
    assert feed.clients_from([a, b])["ABCDE1234F"]["name"] == "Asha Traders Pvt Ltd"
    assert feed.clients_from([b, a])["ABCDE1234F"]["name"] == "Asha Traders Pvt Ltd"
    assert feed.clients_from([a])["ABCDE1234F"]["name"] == "Asha"


def test_csv_write_survives_a_locked_file_and_skips_unchanged_data(tmp_path, monkeypatch):
    x = tmp_path / "LTT_2026-10.xlsx"
    data = _sample()
    assert monthly.write_csvs(x, data)
    data_csv, over_csv = monthly.csv_paths(x)
    stamp = over_csv.read_text(encoding="utf-8-sig")
    assert monthly.write_csvs(x, data, updated=datetime(2030, 1, 1))
    assert over_csv.read_text(encoding="utf-8-sig") == stamp              # unchanged data: nothing rewritten
    real = monthly.os.replace

    def locked(src, dst):
        raise PermissionError(5, "Access is denied")
    monkeypatch.setattr(monthly.os, "replace", locked)
    changed = {k: v[:1] for k, v in data.items()}
    assert monthly.write_csvs(x, changed) is False                        # Excel holds it: no exception
    assert not list(tmp_path.glob("*.tmp"))
    monkeypatch.setattr(monthly.os, "replace", real)
    assert monthly.write_csvs(x, changed) is True


def test_latest_capture_is_on_top():
    cs = [_c("AAAAA1111A", "Alpha", ("GST (GSTR-3B)", "September (FY 2026-27)", "Filed", "1", "2026-10-02T08:00:00+00:00")),
          _c("BBBBB2222B", "Bravo", ("GST (GSTR-3B)", "September (FY 2026-27)", "Filed", "2", "2026-10-06T08:00:00+00:00")),
          _c("CCCCC3333C", "Charlie", ("GST (GSTR-3B)", "September (FY 2026-27)", "Filed", "3", "2026-10-04T08:00:00+00:00")),
          _c("DDDDD4444D", "Delta", ("GST (GSTR-3B)", "September (FY 2026-27)", "Filed", "4", "2026-10-04T08:00:00+00:00"))]
    names = [r["Name"] for r in monthly.latest_rows(cs)["GSTR-3B"]]
    assert names == ["Bravo", "Charlie", "Delta", "Alpha"]


def test_a_recaptured_entry_is_on_top_even_within_the_same_minute():
    """The sheet shows the capture time to the minute, but the newest capture must still come first
    when another capture landed in that same minute (Zulu sorts last by name, so it used to fall below)."""
    cs = [_c("ZZZZZ9999Z", "Zulu", ("GST (GSTR-3B)", "September (FY 2026-27)", "Submitted & Verified", "A1", "2026-10-08T08:37:40+00:00"),),
          _c("AAAAA1111A", "Alpha", ("GST (GSTR-3B)", "September (FY 2026-27)", "Submitted & Verified", "A2", "2026-10-08T08:37:20+00:00"))]
    assert [r["Name"] for r in monthly.latest_rows(cs)["GSTR-3B"]] == ["Zulu", "Alpha"]


def _real_db(tmp_path):
    import security
    from database import SeraDatabase
    salt = str(tmp_path / "t.salt")
    security.generate_and_save_salt(salt)
    key = security.derive_key_hex("testpass123", security.load_salt(salt))
    return SeraDatabase(str(tmp_path / "m.db"), key, raw_db_path=str(tmp_path / "r.db"))


def test_captures_without_an_arn_never_overwrite_another_forms_entry(tmp_path):
    """"N/A" is not an ARN: GSTR-1 and GSTR-3B captured together (no ARN yet) both reach the
    client's history and the LTT sheet, in the order captured; a re-resolve rebuilds the same."""
    import json
    db = _real_db(tmp_path)
    pan = "ABCPD1234E"

    def cap(form, period, status, arn="N/A"):
        db.insert_tracker_dump(portal=f"GST Portal ({form})", period_label=period, arn_number=arn,
                               capture_method="SGT_live", status=status, pan=pan, filing_type=form,
                               raw_payload_json=json.dumps({"pan": pan}))

    cap("GSTR-1", "June (FY 2026-27)", "Submitted & Verified")
    cap("GSTR-3B", "June (FY 2026-27)", "Submitted & Verified")
    cap("GSTR-1", "September (FY 2026-27)", "Not Submitted")
    cap("GSTR-3B", "September (FY 2026-27)", "Draft")
    cap("GSTR-3B", "September (FY 2026-27)", "Submitted (Not Verified)", arn="AA2709260001234")
    cap("GSTR-3B", "September (FY 2026-27)", "Submitted & Verified", arn="AA2709260001234")   # same filing, re-captured

    def sheet():
        rows = monthly.latest_rows(db.get_srpf_containers(limit=1000, slim=True))
        return {f: [(r["Latest period"], r["Status"]) for r in rs] for f, rs in rows.items()}

    want = {"GSTR-1": [("September (FY 2026-27)", "Not Submitted")],
            "GSTR-3B": [("September (FY 2026-27)", "Submitted & Verified")]}
    assert sheet() == want
    hist = db.get_srpf_containers(limit=1000, slim=True)[0]["filing_history"]
    assert [(feed.entry_form(h), h["period_label"][:4]) for h in hist] == [
        ("GSTR-1", "June"), ("GSTR-3B", "June"), ("GSTR-1", "Sept"), ("GSTR-3B", "Sept"), ("GSTR-3B", "Sept")]
    db.re_resolve_all_tracker_dumps()
    assert sheet() == want


def test_the_same_arn_on_another_form_is_a_separate_filing():
    from sera_db.srpf import _real_arn, _same_filing_form
    assert _real_arn(" n/a ") == "" and _real_arn("AA2709260001234") == "AA2709260001234"
    assert _same_filing_form("Income Tax (ITR)", "Income Tax (ITR-4)")
    assert _same_filing_form("GST Portal (GSTR3B)", "GST Portal (GSTR-3B)")
    assert not _same_filing_form("Income Tax (ITR-3)", "Income Tax (ITR-4)")
    assert not _same_filing_form("GST Portal (GSTR-1)", "GST Portal (GSTR-3B)")


# ── Serious-bug fixes (2026-10-07) ───────────────────────────────────────────────
G27, G29 = "27ABCPD1234E1ZE", "29ABCPD1234E1ZA"      # one PAN, two states (valid check characters)


def _h(form, period, status, at, gstin="", arn="N/A"):
    return {"portal": f"GST Portal ({form})", "period_label": period, "status": status, "arn": arn,
            "created_at": at, "gstin": gstin}


def test_ladder_does_not_promote_returns_that_have_not_gone_in():
    assert feed.ladder("Not yet filed") == "Not Submitted"
    assert feed.ladder("Yet to be filed") == "Not Submitted"
    assert feed.ladder("Submission pending") == "Not Submitted"
    assert feed.ladder("Payment pending") == "Not Submitted"
    assert feed.ladder("Filed - ARN pending") == "Submitted (Not Verified)"
    assert feed.ladder("Filed, e-verification pending") == "Submitted (Not Verified)"
    assert feed.ladder("Filed") == "Submitted & Verified"
    assert feed.ladder("Submitted (Not Verified)") == "Submitted (Not Verified)"


def test_each_gstin_of_a_pan_gets_its_own_row():
    """A filed registration must never hide an unfiled one of the same PAN."""
    c = {"pan": "ABCPD1234E", "gstin": G27, "company_name": "Asha Traders", "filing_history": [
        _h("GSTR-3B", "September (FY 2026-27)", "Submitted & Verified", "2026-10-05T10:00:00+00:00", G27),
        _h("GSTR-3B", "September (FY 2026-27)", "Draft", "2026-10-06T10:00:00+00:00", G29),
        _h("ITR-4", "AY 2026-27", "Draft", "2026-10-06T11:00:00+00:00", G29)]}
    c["filing_history"][2]["portal"] = "Income Tax (ITR-4)"
    d = monthly.latest_rows([c])
    assert sorted((r["GSTIN"], r["Status"]) for r in d["GSTR-3B"]) == [(G27, "Submitted & Verified"), (G29, "Draft")]
    assert [r["GSTIN"] for r in d["ITR-4"]] == [""]                  # income tax is per PAN


def test_filings_from_before_gstins_were_recorded_join_the_only_gstin():
    c = {"pan": "ABCPD1234E", "gstin": G27, "company_name": "Asha", "filing_history": [
        _h("GSTR-3B", "August (FY 2026-27)", "Filed", "2026-09-05T10:00:00+00:00"),            # older entry: no gstin key
        _h("GSTR-3B", "June (FY 2026-27)", "Filed", "2026-10-05T10:00:00+00:00", G27)]}
    rows = monthly.latest_rows([c])["GSTR-3B"]
    assert [(r["GSTIN"], r["Latest period"]) for r in rows] == [(G27, "August (FY 2026-27)")]


def test_a_label_without_a_month_never_outranks_a_real_month():
    c = {"pan": "ABCPD1234E", "gstin": "", "company_name": "Asha", "filing_history": [
        _h("GSTR-3B", "September (FY 2026-27)", "Draft", "2026-10-05T10:00:00+00:00"),
        _h("GSTR-3B", "Status-Due (FY 2026-27)", "Due Date", "2026-10-06T10:00:00+00:00")]}
    assert monthly.latest_rows([c])["GSTR-3B"][0]["Latest period"] == "September (FY 2026-27)"
    assert monthly.period_key("AY 2026-27", "ITR-4") == date(2026, 3, 1)          # a year-only label is fine for annual forms
    assert monthly.period_key("FY 2025-26", "GSTR-9") == date(2026, 3, 1)


def test_valid_gstin_rejects_a_misread_check_character():
    from sera_db.srpf import valid_gstin
    assert valid_gstin(G27.lower()) == G27
    assert valid_gstin(G27[:-1] + "F") == ""


def test_a_build_that_fails_in_excel_leaves_no_workbook_behind(tmp_path, monkeypatch):
    """Without the swap-in-at-the-end, the half-built shell looked current and was never rebuilt."""
    x = tmp_path / "LTT_2026-10.xlsx"
    forms = list(_sample())

    def broken(*a, **k):
        raise OSError("Excel went away")
    monkeypatch.setattr(monthly, "attach_queries", broken)
    try:
        monthly.ensure_workbook(x, date(2026, 10, 1), forms, excel=True)
    except OSError:
        pass
    assert not x.exists() and not list(tmp_path.glob("*.xlsx"))
    monkeypatch.setattr(monthly, "attach_queries", lambda *a, **k: None)
    monthly.ensure_workbook(x, date(2026, 10, 1), forms, excel=True)
    assert monthly.workbook_matches(x, forms)


def test_gst_sheets_show_the_gstin_and_others_do_not():
    assert "GSTIN" in monthly.sheet_columns("GSTR-3B") and "GSTIN" in monthly.sheet_columns("CMP-08")
    assert "GSTIN" not in monthly.sheet_columns("ITR-4")


def test_a_locked_summary_is_rewritten_once_free_even_if_the_data_did_not_change(tmp_path, monkeypatch):
    x = tmp_path / "LTT_2026-10.xlsx"
    data = _sample()
    assert monthly.write_csvs(x, data)
    data_csv, over_csv = monthly.csv_paths(x)
    real = monthly.os.replace

    def summary_locked(src, dst):
        if Path(dst) == over_csv:
            raise PermissionError(5, "Access is denied")
        return real(src, dst)
    monkeypatch.setattr(monthly.os, "replace", summary_locked)
    changed = {k: v[:1] for k, v in data.items()}
    assert monthly.write_csvs(x, changed) is False
    monkeypatch.setattr(monthly.os, "replace", real)
    assert monthly.write_csvs(x, changed) is True
    import csv
    over = {r["Form"]: r for r in csv.DictReader(open(over_csv, encoding="utf-8-sig"))}
    assert over["GSTR-3B"]["Clients"] == "1"


def test_the_menu_and_the_capture_timer_never_export_at_once(tmp_path, monkeypatch):
    import threading
    import time
    (tmp_path / monthly.FOLDER).mkdir()
    inside, overlap = [0], []

    def slow(*a, **k):
        inside[0] += 1
        overlap.append(inside[0])
        time.sleep(0.2)
        inside[0] -= 1
    monkeypatch.setattr(monthly, "_export_month", slow)
    ts = [threading.Thread(target=monthly.export_month, args=(None, tmp_path)) for _ in range(3)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    assert overlap == [1, 1, 1]


# ---------------------------------------------------------------------------------------------
# An entry whose status changed must reach the sheet: from another PC (sync), and past a CSV that
# Excel holds open. Reported 2026-10-07: "LTT does not update an entry when it is submitted or its
# status is changed".
# ---------------------------------------------------------------------------------------------

def _filing(db, status, arn="N/A", pan="ABCPD1234E", form="GSTR-1", period="September (FY 2026-27)"):
    import json
    from core.dataset_key import compute_dataset_key
    key = compute_dataset_key("GST Portal", "19ABCPD1234E1ZB", form, period)
    return db.insert_tracker_dump(
        portal=f"GST Portal ({form})", period_label=period, arn_number=arn, capture_method="SGT_live",
        status=status, pan=pan, filing_type=form, dataset_key=key, session_id="s1", captured_by="PC1",
        raw_payload_json=json.dumps({"pan": pan, "gstin": "19ABCPD1234E1ZB", "dataset_key": key}))


def _sheet(db):
    rows = monthly.latest_rows(db.get_srpf_containers(limit=1000, slim=True))
    return [(f, r["Status"], r["ARN"]) for f, rs in rows.items() for r in rs]


def _wait(timer):
    if timer is not None:
        timer.join(15)


def test_a_status_changed_on_another_pc_reaches_the_sheet(tmp_path):
    """Sync writes tracker_dump directly; the containers the sheet reads are a local cache built
    from it. Without a rebuild the sheet keeps the old status."""
    db = _real_db(tmp_path)
    _filing(db, "Draft")
    assert _sheet(db) == [("GSTR-1", "Draft", "")]
    with db._connect_raw() as c:                     # what sync_apply does when PC2 files the return
        c.execute("UPDATE tracker_dump SET status = 'Submitted & Verified', arn_number = 'AA190926191296I'")
        c.commit()
    assert _sheet(db) == [("GSTR-1", "Draft", "")]   # the cache is stale: this is the bug
    _wait(monthly.refresh_after_sync(db, delay=0))
    assert _sheet(db) == [("GSTR-1", "Submitted & Verified", "AA190926191296I")]


def test_a_filing_from_another_pc_reaches_the_csv_and_the_windows_are_told_after(tmp_path, monkeypatch):
    import csv
    db = _real_db(tmp_path)
    _filing(db, "Draft")
    (tmp_path / monthly.FOLDER).mkdir()
    monkeypatch.setattr(monthly, "ensure_workbook", lambda xlsx, month, forms, excel=True: xlsx)
    with db._connect_raw() as c:
        c.execute("UPDATE tracker_dump SET status = 'Submitted & Verified', arn_number = 'AA190926191296I'")
        c.commit()
    order = []
    real = monthly.export_month
    monkeypatch.setattr(monthly, "export_month", lambda *a, **k: (order.append("export"), real(*a, **k))[1])
    _wait(monthly.refresh_after_sync(db, delay=0, on_done=lambda: order.append("windows")))
    assert order == ["export", "windows"]            # the windows refresh with the rebuilt data
    data = Path(db.app_dir) / monthly.FOLDER
    text = "".join(open(p, encoding="utf-8-sig").read() for p in data.glob("*.csv") if "overview" not in p.name)
    assert "Submitted & Verified" in text and "AA190926191296I" in text and "Draft" not in text


def test_a_burst_of_sync_batches_makes_one_refresh(tmp_path):
    db = _real_db(tmp_path)
    runs = []
    db.re_resolve_all_tracker_dumps = lambda: runs.append(1)
    timers = [monthly.refresh_after_sync(db, delay=0.3) for _ in range(5)]
    _wait(timers[-1])
    import time
    time.sleep(0.3)
    assert len(runs) == 1


def test_a_csv_held_open_by_excel_is_tried_again(tmp_path, monkeypatch):
    """The CSV was locked past the 6 s of _replace_file when the status changed: the export used
    to give up until some later capture. It retries."""
    import time
    db = _real_db(tmp_path)
    _filing(db, "Submitted & Verified", arn="AA190926191296I")
    (tmp_path / monthly.FOLDER).mkdir()
    monkeypatch.setattr(monthly, "ensure_workbook", lambda xlsx, month, forms, excel=True: xlsx)
    tries = []
    real = monthly.write_csvs

    def flaky(xlsx, data, updated=None):
        tries.append(1)
        return False if len(tries) < 3 else real(xlsx, data, updated)       # locked twice, then free
    monkeypatch.setattr(monthly, "write_csvs", flaky)
    monthly._busy.clear()
    _wait(monthly.schedule_export(db, delay=0, retry_after=0))
    deadline = time.time() + 15
    while len(tries) < 3 and time.time() < deadline:
        time.sleep(0.05)
        _wait(monthly._timer)
    assert len(tries) == 3 and not monthly._busy.is_set()


def test_a_csv_that_stays_locked_stops_retrying(tmp_path, monkeypatch):
    import time
    db = _real_db(tmp_path)
    _filing(db, "Draft")
    (tmp_path / monthly.FOLDER).mkdir()
    monkeypatch.setattr(monthly, "ensure_workbook", lambda xlsx, month, forms, excel=True: xlsx)
    tries = []
    monkeypatch.setattr(monthly, "write_csvs", lambda *a, **k: (tries.append(1), False)[1])
    _wait(monthly.schedule_export(db, delay=0, retry_after=0))
    deadline = time.time() + 15
    while time.time() < deadline and (monthly._timer is not None and monthly._timer.is_alive()):
        _wait(monthly._timer)
    time.sleep(0.2)
    assert len(tries) == monthly.MAX_RETRIES + 1     # the first try and the retries, then it stops
    monthly._busy.clear()
