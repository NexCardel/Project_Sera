"""core/drs/export.py - DRS Excel workbook and CSV live feed generator.

Generates monthly DRS workbooks (app_dir/DRS/DRS_YYYY-MM.xlsx) and companion CSVs.
Matches the exact visual styling, palette, fonts, and Power Query architecture of LTT.
Supports per-form dynamic column layouts, whole MCL roster feeds, and LTT-style tracker extractions.
Refreshes ONLY when new entries appear in tracker dump (debounced event trigger).
"""
from __future__ import annotations

import calendar
import codecs
import io
import os
import re
import threading
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from core.ltt.feed import LEVELS, NOT_APPLICABLE
from core.ltt.monthly import _replace_file, _rgb, period_key

from .config import (
    get_active_columns,
    get_active_columns_by_form,
    get_active_columns_for_form,
    load_drs_config,
)
from .engine import NO_CAPTURES, build_drs_dataset

FOLDER = "DRS"
FORMAT = "DRS live v3"            # stamped in the workbook; any other stamp means an older layout: rebuild
MAX_ROWS = 20000                   # cap for how far down a live sheet's conditional colouring reaches
STYLE_BUFFER_ROWS = 2000           # rows of colouring kept below the current data so growth stays styled

HEAD_FILL = "FF1F6F4A"            # evergreen, the Sera colour
BAND_FILL = "FFF4F8F5"
OVERVIEW_FIELDS = ("Form", "Clients", "Latest period", *LEVELS, NO_CAPTURES, "Updated")
OVERVIEW_WIDTHS = (24, 12, 28, 20, 12, 28, 24, 20, 18)

STATUS_FILL = {
    LEVELS[0]: "FFF4A99B",
    LEVELS[1]: "FFFFDB70",
    LEVELS[2]: "FFC9E06F",
    LEVELS[3]: "FF86D4A0",
    NO_CAPTURES: "FFDDE7F0",
    NOT_APPLICABLE: "FFDADADA",
}

STATUS_TEXT = {
    LEVELS[0]: "FF7A1F12",
    LEVELS[1]: "FF6B4A00",
    LEVELS[2]: "FF4A5A0A",
    LEVELS[3]: "FF0F4F2A",
    NO_CAPTURES: "FF35506B",
    NOT_APPLICABLE: "FF555555",
}

DEFAULT_WIDTHS = {
    "PAN": 14,
    "GSTIN": 19,
    "Name": 36,
    "Client Name": 36,
    "Company Name": 36,
    "Trade Name": 36,
    "Latest period": 26,
    "Status": 24,
    "ARN": 24,
    "Captured on": 18,
    "Captured by": 18,
    "Capture method": 18,
    "Original author": 18,
    "Portal": 18,
    "Portal Status": 18,
    "Filing Period": 24,
    "Portal ID": 18,
    "Dataset Key": 28,
    "Notes": 32,
    "Mobile": 16,
    "Phone": 16,
    "Email": 28,
}


def month_start(d: date) -> date:
    return d.replace(day=1)


def drs_workbook_path(app_dir, month: date) -> Path:
    return Path(app_dir) / FOLDER / f"DRS_{month:%Y-%m}.xlsx"


def drs_csv_paths(xlsx: Path) -> tuple[Path, Path]:
    x = Path(xlsx)
    return x.with_suffix(".csv"), x.with_name(x.stem + "_overview.csv")


def _latest_period(rows: list[dict], form: str) -> str:
    """The newest period actually present on a sheet, not just the first row's period."""
    best_label = NO_CAPTURES
    best_key = None
    fallback = NO_CAPTURES
    for row in rows:
        label = str(row.get("Latest period") or "").strip()
        if not label or label == NO_CAPTURES:
            continue
        if fallback == NO_CAPTURES:
            fallback = label
        key = period_key(label, form)
        if key is not None and (best_key is None or key > best_key):
            best_key, best_label = key, label
    return best_label if best_key is not None else fallback


def _open_in_excel(xlsx: Path) -> bool:
    try:
        with open(xlsx, "r+b"):
            return False
    except FileNotFoundError:
        return False
    except OSError:
        return True


def _has_live_tables(xlsx: Path) -> bool:
    """True when Excel COM has attached Power Query tables to this workbook."""
    try:
        with zipfile.ZipFile(xlsx) as z:
            return any(name.startswith(("xl/tables/", "xl/queryTables/", "xl/connections"))
                       for name in z.namelist())
    except Exception:
        return False


def refresh_open_drs_workbook(xlsx: Path) -> bool:
    """Refreshes the live queries in the workbook the user already has open."""
    try:
        import comtypes
        import comtypes.client

        comtypes.CoInitialize()
        xl = comtypes.client.GetActiveObject("Excel.Application")
        target = os.path.normcase(os.path.abspath(xlsx))
        for wb in xl.Workbooks:
            if os.path.normcase(os.path.abspath(wb.FullName)) != target:
                continue
            wb.RefreshAll()
            return True
    except Exception:
        pass
    return False


def _normalize_cols_map(cols_arg: Any, form_list: list[str]) -> dict[str, list[dict[str, Any]]]:
    if isinstance(cols_arg, dict):
        return cols_arg
    if isinstance(cols_arg, list):
        return {f: cols_arg for f in form_list}
    return {}


def workbook_matches(xlsx: Path, forms: list[str],
                     active_cols_by_form: dict[str, list[dict[str, Any]]] | list[dict[str, Any]]) -> bool:
    """Checks that every current form tab and its specific column headers match."""
    wb = None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, read_only=True)
        if wb.properties.title != FORMAT:
            return False
        expected = dict(sheet_names(forms))
        expected_names = {"Overview", *expected.values()}
        if set(wb.sheetnames) != expected_names:
            return False

        cols_map = _normalize_cols_map(active_cols_by_form, forms)
        for form, sheet in expected.items():
            ws = wb[sheet]
            cols = cols_map.get(form, [])
            labels = [c["label"] for c in cols]
            headers = [ws.cell(row=3, column=i + 1).value for i in range(len(labels))]
            if [str(v or "") for v in headers] != labels:
                return False
        return True
    except Exception:
        return False
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def _csv_bytes(fields, rows) -> bytes:
    """UTF-8 with BOM for Excel compatibility."""
    import csv
    buf = io.StringIO(newline="")
    w = csv.DictWriter(buf, fieldnames=fields, extrasaction="ignore")
    w.writeheader()
    w.writerows(rows)
    return codecs.BOM_UTF8 + buf.getvalue().encode("utf-8")


def _sheet_name(form: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", form).strip()[:31] or "Form"
    name, i = base, 2
    while name.lower() in used or name.lower() == "overview":
        suffix = f" ({i})"
        name, i = base[:31 - len(suffix)] + suffix, i + 1
    used.add(name.lower())
    return name


def sheet_names(forms: list[str]) -> dict[str, str]:
    used: set = set()
    return {f: _sheet_name(f, used) for f in forms}


def build_drs_shell(path: Path, month: date, forms: list[str],
                    active_cols_by_form: dict[str, list[dict[str, Any]]] | list[dict[str, Any]]) -> Path:
    """Builds the workbook frame with the same layout, palette and print setup as LTT."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.properties.title = FORMAT

    # Overview sheet
    ov = wb.active
    ov.title = "Overview"

    font_sub = Font(name="Calibri", size=10, color="FF555555")
    font_title = Font(name="Calibri", size=18, bold=True, color=HEAD_FILL)
    font_form_title = Font(name="Calibri", size=16, bold=True, color=HEAD_FILL)

    def frame(ws, widths):
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    frame(ov, OVERVIEW_WIDTHS)
    ov["A1"] = f"DRS — {calendar.month_name[month.month]} {month.year}"
    ov["A1"].font = font_title
    ov["A2"] = "Waiting for data…"
    ov["A2"].font = font_sub
    ov["A3"] = ("Every registered client enrolled in each portal service, with the latest filing period, "
                "status, ARN and capture details. Live while open in Excel.")
    ov["A3"].font = font_sub
    ov.row_dimensions[4].height = 30
    ov.freeze_panes = "A5"
    ov.column_dimensions[get_column_letter(len(OVERVIEW_FIELDS))].hidden = True  # "Updated" is shown in A2

    names = sheet_names(forms)
    cols_map = _normalize_cols_map(active_cols_by_form, forms)

    for form, sname in names.items():
        ws = wb.create_sheet(sname)
        form_cols = cols_map.get(form, [])
        col_labels = [c["label"] for c in form_cols]
        frame(ws, (DEFAULT_WIDTHS.get(lbl, 20) for lbl in col_labels))
        ws["A1"] = form
        ws["A1"].font = font_form_title
        ws["A2"] = "Waiting for data…"
        ws["A2"].font = font_sub
        ws.row_dimensions[3].height = 22
        ws.freeze_panes = "A4"
        ws.print_title_rows = "3:3"

    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)
    return path


def _style_drs_table(ws, lo, r, first_row: int, centre: tuple = (), status_col: str | None = None,
                     mono: tuple = (), data_rows: int | None = None) -> None:
    """Same table styling as LTT: body fonts, alignment, status conditional colours, banding and borders."""
    from openpyxl.utils import get_column_letter
    count = r(lambda: lo.ListColumns.Count)
    last_col = get_column_letter(count)

    # One range-level format pass instead of ~6 COM calls for every column.
    body = r(lambda: lo.DataBodyRange)
    if body is not None:
        f = r(lambda: body.Font)
        r(lambda: setattr(f, "Name", "Calibri"))
        r(lambda: setattr(f, "Size", 10))
        r(lambda: setattr(body, "HorizontalAlignment", -4131))
        r(lambda: setattr(body, "IndentLevel", 1))
        r(lambda: setattr(body, "VerticalAlignment", -4108))
        for i in mono:
            col_body = r(lambda: lo.ListColumns(i).DataBodyRange)
            if col_body is not None:
                r(lambda: setattr(r(lambda: col_body.Font), "Name", "Consolas"))
        for i in centre:
            col_body = r(lambda: lo.ListColumns(i).DataBodyRange)
            if col_body is not None:
                r(lambda: setattr(col_body, "HorizontalAlignment", -4108))
                r(lambda: setattr(col_body, "IndentLevel", 0))

    hdr = r(lambda: lo.HeaderRowRange)
    r(lambda: setattr(hdr, "HorizontalAlignment", -4131))
    r(lambda: setattr(hdr, "IndentLevel", 1))
    r(lambda: ws.Activate())
    r(lambda: ws.Range("A1").Select())

    # Colour only as far as this sheet's data can plausibly reach. Formatting 20k rows on every
    # sheet made Excel slow to open and scroll; keep a generous buffer for later captures.
    if data_rows is None:
        styled_rows = MAX_ROWS
    else:
        styled_rows = min(MAX_ROWS, max(first_row + max(data_rows, 0) + STYLE_BUFFER_ROWS, 1000))

    rng = r(lambda: ws.Range(f"A1:{last_col}{styled_rows}"))
    r(lambda: rng.FormatConditions.Delete())
    body_row = f"ROW()>={first_row}"
    if status_col:
        srng = r(lambda: ws.Range(f"{status_col}1:{status_col}{styled_rows}"))
        for lvl, fill in STATUS_FILL.items():
            fc = r(lambda: srng.FormatConditions.Add(2, 0, f'=AND({body_row},${status_col}1="{lvl}")'))
            r(lambda: setattr(r(lambda: fc.Interior), "Color", _rgb(fill)))
            r(lambda: setattr(r(lambda: fc.Font), "Color", _rgb(STATUS_TEXT.get(lvl, "FF000000"))))
            r(lambda: setattr(r(lambda: fc.Font), "Bold", True))
    row_has_data = f"COUNTA($A1:${last_col}1)>0"
    fc = r(lambda: rng.FormatConditions.Add(2, 0, f"=AND({body_row},{row_has_data},MOD(ROW(),2)=0)"))
    r(lambda: setattr(r(lambda: fc.Interior), "Color", _rgb(BAND_FILL)))
    fc = r(lambda: rng.FormatConditions.Add(2, 0, f"=AND({body_row},{row_has_data})"))
    edge = r(lambda: fc.Borders(-4107))
    r(lambda: setattr(edge, "LineStyle", 1))
    r(lambda: setattr(edge, "Color", _rgb("FFD9D9D9")))


def attach_drs_queries(xlsx: Path, forms: list[str],
                       active_cols_by_form: dict[str, list[dict[str, Any]]] | list[dict[str, Any]],
                       month_file: Path | None = None,
                       data_rows_by_form: dict[str, int] | None = None) -> bool:
    """Attaches Power Query live connections via Windows COM Excel if available.
    `month_file` is the final workbook path when `xlsx` is a temporary build copy.
    `data_rows_by_form` lets conditional colouring stop near the real data instead of far below it."""
    try:
        import comtypes
        import comtypes.client
    except ImportError:
        return False

    data_csv, over_csv = drs_csv_paths(month_file or xlsx)
    if not data_csv.exists() or not over_csv.exists():
        return False

    try:
        comtypes.CoInitialize()
    except Exception:
        pass

    def r(fn, tries=15, pause=0.25):
        import time
        for i in range(tries):
            try:
                return fn()
            except Exception:
                if i < tries - 1:
                    time.sleep(pause)
                else:
                    raise

    from openpyxl.utils import get_column_letter

    names = sheet_names(forms)
    cols_map = _normalize_cols_map(active_cols_by_form, forms)
    esc = lambda p: str(p).replace('"', '""')  # noqa: E731
    head = ('let Source = Csv.Document(File.Contents("%s"), [Delimiter=",", Encoding=65001, QuoteStyle=QuoteStyle.Csv]), '
            'Promoted = Table.PromoteHeaders(Source, [PromoteAllScalars=true])')

    jobs = [(
        "Overview", "A4", "DRS_Overview",
        head % esc(over_csv.resolve())
        + ', Typed = Table.TransformColumnTypes(Promoted, '
          'List.Transform(Table.ColumnNames(Promoted), each {_, type text})) in Typed',
        OVERVIEW_FIELDS,
        len(forms),
    )]

    for i, (form, sname) in enumerate(names.items(), 1):
        cols = cols_map.get(form, [])
        col_labels = [c["label"] for c in cols]
        cols_arg = ", ".join(f'"{c}"' for c in col_labels)
        # Filter and trim the CSV before converting types: converting every column of every
        # form's full scan made each refresh noticeably heavier.
        m = (head % esc(data_csv.resolve())
             + f', Rows = Table.SelectRows(Promoted, each [Form] = "{esc(form)}"), '
               f'Cols = Table.SelectColumns(Rows, {{{cols_arg}}}), '
               f'Typed = Table.TransformColumnTypes(Cols, '
               f'List.Transform(Table.ColumnNames(Cols), each {{_, type text}})) in Typed')
        rows = (data_rows_by_form or {}).get(form, 0)
        jobs.append((sname, "A3", f"DRS_Form_{i}", m, col_labels, rows))

    xl = None
    wb = None
    ok = False
    try:
        xl = comtypes.client.CreateObject("Excel.Application")
        r(lambda: setattr(xl, "Visible", False))
        r(lambda: setattr(xl, "DisplayAlerts", False))
        # Build without repainting or event handling; recalculate once at the end.
        r(lambda: setattr(xl, "ScreenUpdating", False))
        r(lambda: setattr(xl, "EnableEvents", False))
        wb = r(lambda: xl.Workbooks.Open(str(xlsx.resolve())))
        r(lambda: setattr(xl, "Calculation", -4135))      # xlCalculationManual

        # Add every query first and let Power Query refresh them together in the background.
        # Refreshing one table at a time serialised the whole build (and the Query engine's
        # start-up cost) into a long wait.
        tables = []
        for sheet, anchor, query, m, cols, rows in jobs:
            r(lambda: wb.Queries.Add(query, m))
            ws = r(lambda: wb.Worksheets(sheet))
            lo = r(lambda: ws.ListObjects.Add(
                0, 'OLEDB;Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;'
                   f'Location={query};Extended Properties=""', None, 1, r(lambda: ws.Range(anchor))))
            r(lambda: setattr(lo, "Name", query))
            qt = r(lambda: lo.QueryTable)
            r(lambda: setattr(qt, "CommandType", 2))
            r(lambda: setattr(qt, "CommandText", f"SELECT * FROM [{query}]"))
            r(lambda: setattr(qt, "BackgroundQuery", True))
            r(lambda: setattr(qt, "RefreshStyle", 0))
            r(lambda: setattr(qt, "AdjustColumnWidth", False))
            r(lambda: setattr(qt, "PreserveColumnInfo", True))
            r(lambda: setattr(qt, "PreserveFormatting", True))
            r(lambda: qt.Refresh(True))
            tables.append((sheet, query, ws, lo, qt, cols, rows))

        try:
            xl.CalculateUntilAsyncQueriesDone()
        except Exception:
            pass

        for sheet, query, ws, lo, qt, cols, rows in tables:
            # An async refresh that failed would leave an empty table: retry that one synchronously.
            expected = rows
            if expected:
                body = r(lambda: lo.DataBodyRange)
                actual = r(lambda: body.Rows.Count) if body is not None else 0
                if actual != expected:
                    r(lambda: setattr(qt, "BackgroundQuery", False))
                    r(lambda: qt.Refresh(False))
                    r(lambda: setattr(qt, "BackgroundQuery", True))

            r(lambda: setattr(lo, "TableStyle", ""))
            hdr = r(lambda: lo.HeaderRowRange)
            r(lambda: setattr(r(lambda: hdr.Interior), "Color", _rgb(HEAD_FILL)))
            r(lambda: setattr(r(lambda: hdr.Font), "Color", _rgb("FFFFFFFF")))
            r(lambda: setattr(r(lambda: hdr.Font), "Bold", True))
            r(lambda: setattr(hdr, "VerticalAlignment", -4108))

            if sheet == "Overview":
                _style_drs_table(ws, lo, r, 5, centre=(2, 4, 5, 6, 7, 8), data_rows=rows)
                f = (f'=IF(INDEX({query}[Updated],1)="","Waiting for the first capture…",'
                     f'"Updated "&INDEX({query}[Updated],1))')
            else:
                labels = list(cols or [])
                status_col = get_column_letter(labels.index("Status") + 1) if "Status" in labels else None
                mono = tuple(i + 1 for i, c in enumerate(labels) if c in ("PAN", "GSTIN", "ARN", "Portal ID", "Dataset Key"))
                _style_drs_table(ws, lo, r, 4, centre=(), status_col=status_col, mono=mono, data_rows=rows)
                parts = [f'ROWS({query}[#Data])&" clients"']
                if "Status" in labels:
                    for lvl in (*LEVELS, NO_CAPTURES, NOT_APPLICABLE):
                        n = f'COUNTIF({query}[Status],"{lvl}")'
                        parts.append(f'IF({n}>0,"  ·  "&{n}&" {lvl}","")')
                f = "=" + "&".join(parts)
            r(lambda: setattr(r(lambda: ws.Range("A2")), "Formula", f))

            conn = r(lambda: qt.WorkbookConnection.OLEDBConnection)
            r(lambda: setattr(conn, "BackgroundQuery", True))
            r(lambda: setattr(conn, "RefreshOnFileOpen", True))
            r(lambda: setattr(conn, "RefreshPeriod", 0))

        r(lambda: setattr(xl, "Calculation", -4105))      # xlCalculationAutomatic
        r(lambda: xl.CalculateFull())
        r(lambda: setattr(xl, "ScreenUpdating", True))
        r(lambda: setattr(xl, "EnableEvents", True))
        r(lambda: wb.Worksheets("Overview").Activate())
        r(lambda: wb.Save())
        ok = True
        r(lambda: wb.Close(False))
        wb = None
    except Exception as e:
        print(f"[drs] Power Query attachment fallback: {e}")
    finally:
        if wb is not None:
            try: r(lambda: wb.Close(False), tries=10)
            except Exception: pass
        if xl is not None:
            try: r(lambda: xl.Quit(), tries=10)
            except Exception: pass
    return ok


def populate_direct_openpyxl(xlsx: Path, month: date, data: dict[str, list[dict]],
                             active_cols_by_form: dict[str, list[dict[str, Any]]] | list[dict[str, Any]],
                             updated: datetime) -> None:
    """Fills data directly into openpyxl sheets when Excel COM / live Power Query is absent."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = load_workbook(xlsx)
    names = sheet_names(list(data.keys()))
    cols_map = _normalize_cols_map(active_cols_by_form, list(names.keys()))

    font_head = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    font_body = Font(name="Calibri", size=10, color="FF000000")
    font_mono = Font(name="Consolas", size=10, color="FF000000")
    fill_head = PatternFill(start_color=HEAD_FILL, end_color=HEAD_FILL, fill_type="solid")
    fill_band = PatternFill(start_color=BAND_FILL, end_color=BAND_FILL, fill_type="solid")
    align_left = Alignment(horizontal="left", vertical="center", indent=1)
    align_center = Alignment(horizontal="center", vertical="center")
    edge = Border(bottom=Side(style="thin", color="FFD9D9D9"))

    # build each cell style once and reuse its StyleArray: assigning font/fill/border/alignment
    # cell-by-cell makes openpyxl hash a fresh style object for every cell and dominated export time.
    tpl_name = "__drs_styles"
    while tpl_name in wb.sheetnames:
        tpl_name += "_"
    tpl = wb.create_sheet(tpl_name)
    tpl_row = 0

    def _style(font, fill=None, alignment=align_left, border=None):
        nonlocal tpl_row
        tpl_row += 1
        cell = tpl.cell(row=tpl_row, column=1)
        cell.font = font
        cell.alignment = alignment
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
        return cell._style

    head_style = _style(font_head, fill_head)
    body_left = {False: _style(font_body, None, align_left, edge),
                 True: _style(font_body, fill_band, align_left, edge)}
    body_mono = {False: _style(font_mono, None, align_left, edge),
                 True: _style(font_mono, fill_band, align_left, edge)}
    body_center = {False: _style(font_body, None, align_center, edge),
                   True: _style(font_body, fill_band, align_center, edge)}
    status_styles = {
        lvl: _style(Font(name="Calibri", size=10, bold=True, color=STATUS_TEXT.get(lvl, "FF000000")),
                    PatternFill(start_color=fill, end_color=fill, fill_type="solid"), align_left, edge)
        for lvl, fill in STATUS_FILL.items()
    }

    # 1. Update Overview Sheet
    if "Overview" in wb.sheetnames:
        ov = wb["Overview"]
        ov["A2"] = f"Updated {updated:%d %b %Y %H:%M}"
        if ov.max_row >= 5:
            ov.delete_rows(5, ov.max_row - 4)

        for c_idx, h in enumerate(OVERVIEW_FIELDS, 1):
            cell = ov.cell(row=4, column=c_idx, value=h)
            cell._style = head_style
        ov.column_dimensions[get_column_letter(len(OVERVIEW_FIELDS))].hidden = True

        r_idx = 5
        for form, rows in data.items():
            counts = {lvl: sum(1 for r in rows if r.get("Status") == lvl) for lvl in (*LEVELS, NO_CAPTURES)}
            best_period = _latest_period(rows, form)
            row_vals = [form, len(rows), best_period] + [counts.get(lvl, 0) for lvl in (*LEVELS, NO_CAPTURES)] + [f"{updated:%d %b %Y %H:%M}"]
            band = r_idx % 2 == 0
            for c_idx, v in enumerate(row_vals, 1):
                cell = ov.cell(row=r_idx, column=c_idx, value=v)
                cell._style = (body_center if c_idx in (2, 4, 5, 6, 7, 8) else body_left)[band]
            r_idx += 1
        ov.auto_filter.ref = f"A4:{get_column_letter(len(OVERVIEW_FIELDS))}{max(4, r_idx - 1)}"

    # 2. Update Form Sheets
    for form, sname in names.items():
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        rows = data.get(form, [])
        cols = cols_map.get(form, [])
        col_labels = [c["label"] for c in cols]
        last_col = get_column_letter(len(col_labels)) if col_labels else "A"

        parts = [f"{len(rows)} clients"]
        if "Status" in col_labels:
            for lvl in (*LEVELS, NO_CAPTURES, NOT_APPLICABLE):
                n = sum(1 for r in rows if r.get("Status") == lvl)
                if n:
                    parts.append(f"{n} {lvl}")
        ws["A2"] = "  ·  ".join(parts)
        ws.print_title_rows = "3:3"

        if ws.max_row >= 4:
            ws.delete_rows(4, ws.max_row - 3)
        for c_idx in range(len(col_labels) + 1, ws.max_column + 1):
            ws.cell(row=3, column=c_idx, value=None)

        # Headers at row 3
        for c_idx, lbl in enumerate(col_labels, 1):
            cell = ws.cell(row=3, column=c_idx, value=lbl)
            cell._style = head_style

        mono_cols = [any(kw in lbl.upper() for kw in ("PAN", "GSTIN", "ARN", "PORTAL ID", "DATASET KEY"))
                     for lbl in col_labels]
        status_idx = col_labels.index("Status") if "Status" in col_labels else -1

        # Rows starting at row 4
        for r_idx, r_data in enumerate(rows, 4):
            band = r_idx % 2 == 0
            for c_idx, lbl in enumerate(col_labels, 1):
                val = r_data.get(lbl, "")
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                # Status colouring, matching live table conditional formatting
                if c_idx == status_idx + 1 and val in status_styles:
                    cell._style = status_styles[val]
                else:
                    cell._style = (body_mono if mono_cols[c_idx - 1] else body_left)[band]

        if col_labels:
            ws.auto_filter.ref = f"A3:{last_col}{max(3, len(rows) + 3)}"

    del wb[tpl_name]
    wb.save(xlsx)


def write_drs_csvs(xlsx: Path, data: dict[str, list[dict]],
                   active_cols_by_form: dict[str, list[dict[str, Any]]] | list[dict[str, Any]],
                   updated: datetime) -> bool:
    """Writes DRS companion CSV files supporting per-form column structures."""
    data_csv, over_csv = drs_csv_paths(xlsx)
    cols_map = _normalize_cols_map(active_cols_by_form, list(data.keys()))

    # Gather union of column labels across all forms, keeping original sequence
    seen_labels = set()
    ordered_labels = []
    for form in data.keys():
        for c in cols_map.get(form, []):
            lbl = c["label"]
            if lbl not in seen_labels:
                seen_labels.add(lbl)
                ordered_labels.append(lbl)

    flat_data = []
    for form, rows in data.items():
        for r in rows:
            entry = {"Form": form}
            for lbl in ordered_labels:
                entry[lbl] = r.get(lbl, "")
            flat_data.append(entry)

    over_rows = []
    for form, rows in data.items():
        counts = {lvl: sum(1 for r in rows if r.get("Status") == lvl) for lvl in (*LEVELS, NO_CAPTURES)}
        best_p = _latest_period(rows, form)
        over_rows.append({
            "Form": form,
            "Clients": len(rows),
            "Latest period": best_p,
            **counts,
            "Updated": f"{updated:%d %b %Y %H:%M}",
        })

    over_bytes = _csv_bytes(OVERVIEW_FIELDS, over_rows)
    data_bytes = _csv_bytes(["Form"] + ordered_labels, flat_data)

    return _replace_file(over_csv, over_bytes) and _replace_file(data_csv, data_bytes)


_export_lock = threading.Lock()


def export_drs(db, app_dir=None, today: date | None = None, force_rebuild: bool = False,
               excel: bool = True) -> Path | None:
    """Builds and refreshes the current month's DRS workbook and CSVs.
    Returns the path to the .xlsx workbook.
    """
    today = today or date.today()
    app_dir = Path(app_dir or db.app_dir)
    month = month_start(today)
    xlsx = drs_workbook_path(app_dir, month)
    xlsx.parent.mkdir(parents=True, exist_ok=True)

    with _export_lock:
        cfg = load_drs_config(db)
        data = build_drs_dataset(db, config=cfg)
        forms = list(data.keys())
        active_cols_by_form = get_active_columns_by_form(db, forms)
        updated = datetime.now(timezone.utc).astimezone()

        # 1. Companion feeds must exist before Excel COM reads them
        written = write_drs_csvs(xlsx, data, active_cols_by_form, updated)

        # 2. Check if workbook already matches
        matches = xlsx.exists() and not force_rebuild and workbook_matches(xlsx, forms, active_cols_by_form)
        if matches:
            if _has_live_tables(xlsx):
                if written and _open_in_excel(xlsx):
                    refresh_open_drs_workbook(xlsx)
                return xlsx
            if not excel:
                populate_direct_openpyxl(xlsx, month, data, active_cols_by_form, updated)
                return xlsx
            if _open_in_excel(xlsx):
                return xlsx

        if _open_in_excel(xlsx):
            raise RuntimeError(f"{xlsx.name} is open in Excel and needs rebuilding - close it and try again.")

        # 3. Build the shell aside and swap it in once complete
        work = xlsx.with_name(xlsx.stem + ".building.xlsx")
        try:
            build_drs_shell(work, month, forms, active_cols_by_form)
            live = (attach_drs_queries(work, forms, active_cols_by_form, month_file=xlsx,
                                       data_rows_by_form={f: len(rows) for f, rows in data.items()})
                    if excel else False)
            if not live:
                if excel:
                    # COM may have half-attached queries before failing: start from a clean shell.
                    build_drs_shell(work, month, forms, active_cols_by_form)
                populate_direct_openpyxl(work, month, data, active_cols_by_form, updated)
            os.replace(work, xlsx)
        finally:
            try:
                work.unlink(missing_ok=True)
            except OSError:
                pass

        return xlsx


_timer = None
_timer_lock = threading.Lock()


def schedule_drs_export(db, delay: float = 3.0):
    """Event-driven debounced export trigger.
    Refreshes ONLY when called upon new capture or sync arrival.
    Does NOT poll on fixed intervals.
    """
    global _timer
    try:
        def run():
            try:
                export_drs(db)
            except Exception as e:
                print(f"[drs] scheduled refresh skipped: {e}")

        with _timer_lock:
            if _timer is not None:
                _timer.cancel()
            _timer = threading.Timer(delay, run)
            _timer.daemon = True
            _timer.start()
            return _timer
    except Exception:
        return None


def open_drs_workbook(app_dir: str | Path, today: date | None = None) -> bool:
    """Opens the DRS workbook in Excel or default spreadsheet handler."""
    today = today or date.today()
    xlsx = drs_workbook_path(app_dir, month_start(today))
    if not xlsx.exists():
        return False
    try:
        os.startfile(str(xlsx))
        return True
    except Exception as e:
        print(f"[drs] Failed to open workbook: {e}")
        return False
